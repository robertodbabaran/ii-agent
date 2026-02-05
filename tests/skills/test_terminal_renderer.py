"""
Unit tests for terminal_renderer.py — ModelMetrics, extract_metrics, TerminalModelRenderer.
"""

import pytest
from pathlib import Path
from unittest.mock import patch, MagicMock

from ii_skills.shared.terminal_renderer import (
    ModelMetrics,
    TerminalModelRenderer,
    extract_metrics,
    render_from_assumptions,
    _safe_float,
    _fmt_dollar,
    _fmt_pct,
    _fmt_mult,
    _compute_sources_uses,
    _compute_projections,
    _compute_debt_schedule,
    _compute_returns,
    _compute_sensitivity,
)


# ---------------------------------------------------------------------------
# Helper formatting
# ---------------------------------------------------------------------------

class TestFormatHelpers:

    def test_safe_float_none(self):
        assert _safe_float(None) == 0.0

    def test_safe_float_int(self):
        assert _safe_float(5) == 5.0

    def test_safe_float_float(self):
        assert _safe_float(3.14) == 3.14

    def test_safe_float_string(self):
        assert _safe_float("1,234.5") == 1234.5

    def test_safe_float_dollar_string(self):
        assert _safe_float("$100") == 100.0

    def test_safe_float_formula(self):
        assert _safe_float("=B5*B6") == 0.0

    def test_safe_float_garbage(self):
        assert _safe_float("abc") == 0.0

    def test_fmt_dollar_millions(self):
        assert _fmt_dollar(123.4) == "$123.4M"

    def test_fmt_dollar_billions(self):
        assert _fmt_dollar(1500.0) == "$1.5B"

    def test_fmt_pct(self):
        assert _fmt_pct(0.25) == "25.0%"

    def test_fmt_pct_zero(self):
        assert _fmt_pct(0.0) == "0.0%"

    def test_fmt_mult(self):
        assert _fmt_mult(8.0) == "8.0x"


# ---------------------------------------------------------------------------
# ModelMetrics computation
# ---------------------------------------------------------------------------

class TestModelMetricsComputation:

    def _make_metrics(self) -> ModelMetrics:
        m = ModelMetrics()
        m.company_name = "Test Corp"
        m.ltm_revenue = 100.0
        m.ltm_ebitda = 20.0
        m.entry_multiple = 8.0
        m.exit_multiple = 8.0
        m.hold_period = 5
        m.senior_debt_multiple = 4.0
        m.senior_interest_rate = 0.08
        m.senior_amortization = 0.01
        m.sub_debt_multiple = 1.5
        m.sub_interest_rate = 0.12
        m.transaction_fee_pct = 0.015
        m.financing_fee_pct = 0.025
        m.revenue_growth = [0.08, 0.07, 0.06, 0.05, 0.05]
        m.ebitda_margin = [0.20, 0.21, 0.22, 0.22, 0.22]
        return m

    def test_sources_uses(self):
        m = self._make_metrics()
        _compute_sources_uses(m)

        assert m.entry_ev == pytest.approx(160.0)  # 20 * 8
        assert m.total_senior_debt == pytest.approx(80.0)  # 20 * 4
        assert m.total_sub_debt == pytest.approx(30.0)  # 20 * 1.5
        assert m.total_debt == pytest.approx(110.0)
        assert m.transaction_fees == pytest.approx(2.4)  # 160 * 0.015
        assert m.financing_fees == pytest.approx(2.75)  # 110 * 0.025
        assert m.equity_invested == pytest.approx(55.15)  # 160 - 110 + 2.4 + 2.75

    def test_projections(self):
        m = self._make_metrics()
        _compute_sources_uses(m)
        _compute_projections(m)

        assert len(m.projected_revenue) == 5
        assert len(m.projected_ebitda) == 5

        # Year 1: 100 * 1.08 = 108.0
        assert m.projected_revenue[0] == pytest.approx(108.0)
        # Year 1 EBITDA: 108.0 * 0.20 = 21.6
        assert m.projected_ebitda[0] == pytest.approx(21.6)

    def test_debt_schedule(self):
        m = self._make_metrics()
        _compute_sources_uses(m)
        _compute_projections(m)
        _compute_debt_schedule(m)

        assert len(m.total_debt_by_year) == 5
        assert len(m.leverage_by_year) == 5

        # Debt should decrease over time due to amortization
        assert m.total_debt_by_year[0] < m.total_debt

    def test_returns(self):
        m = self._make_metrics()
        _compute_sources_uses(m)
        _compute_projections(m)
        _compute_debt_schedule(m)
        _compute_returns(m)

        assert m.exit_ebitda > 0
        assert m.exit_ev > 0
        assert m.moic > 0
        assert m.irr > 0

    def test_sensitivity(self):
        m = self._make_metrics()
        _compute_sources_uses(m)
        _compute_projections(m)
        _compute_debt_schedule(m)
        _compute_returns(m)
        _compute_sensitivity(m)

        assert len(m.sensitivity_rows) > 0
        assert len(m.sensitivity_entry_range) == 5
        assert len(m.sensitivity_exit_range) == 5

        # Each row is (entry_mult, exit_mult, moic, irr)
        for row in m.sensitivity_rows:
            assert len(row) == 4


# ---------------------------------------------------------------------------
# TerminalModelRenderer — text rendering
# ---------------------------------------------------------------------------

class TestTerminalModelRendering:

    def _make_full_metrics(self) -> ModelMetrics:
        m = ModelMetrics()
        m.company_name = "Acme Corp"
        m.ltm_revenue = 100.0
        m.ltm_ebitda = 20.0
        m.entry_multiple = 8.0
        m.exit_multiple = 8.0
        m.hold_period = 5
        m.senior_debt_multiple = 4.0
        m.senior_interest_rate = 0.08
        m.senior_amortization = 0.01
        m.sub_debt_multiple = 1.5
        m.sub_interest_rate = 0.12
        m.transaction_fee_pct = 0.015
        m.financing_fee_pct = 0.025
        m.revenue_growth = [0.08, 0.07, 0.06, 0.05, 0.05]
        m.ebitda_margin = [0.20, 0.21, 0.22, 0.22, 0.22]
        m.sheets_found = ["Assumptions", "Sources & Uses", "Operating Model"]

        _compute_sources_uses(m)
        _compute_projections(m)
        _compute_debt_schedule(m)
        _compute_returns(m)
        _compute_sensitivity(m)
        return m

    def test_render_deal_summary(self):
        m = self._make_full_metrics()
        renderer = TerminalModelRenderer(use_color=False)
        text = renderer.render_deal_summary(m)

        assert "Acme Corp" in text
        assert "8.0x" in text
        assert "MOIC" in text
        assert "IRR" in text

    def test_render_sources_uses(self):
        m = self._make_full_metrics()
        renderer = TerminalModelRenderer(use_color=False)
        text = renderer.render_sources_uses(m)

        assert "SOURCES" in text
        assert "USES" in text
        assert "Senior Debt" in text
        assert "Sponsor Equity" in text

    def test_render_operating_model(self):
        m = self._make_full_metrics()
        renderer = TerminalModelRenderer(use_color=False)
        text = renderer.render_operating_model(m)

        assert "Revenue" in text
        assert "EBITDA" in text
        assert "Year 1" in text
        assert "Year 5" in text

    def test_render_debt_summary(self):
        m = self._make_full_metrics()
        renderer = TerminalModelRenderer(use_color=False)
        text = renderer.render_debt_summary(m)

        assert "Total Debt" in text
        assert "Leverage" in text

    def test_render_returns(self):
        m = self._make_full_metrics()
        renderer = TerminalModelRenderer(use_color=False)
        text = renderer.render_returns(m)

        assert "Equity Invested" in text
        assert "MOIC" in text
        assert "IRR" in text
        assert "Exit EV" in text

    def test_render_sensitivity(self):
        m = self._make_full_metrics()
        renderer = TerminalModelRenderer(use_color=False)
        text = renderer.render_sensitivity(m)

        assert "MOIC Sensitivity" in text
        assert "IRR Sensitivity" in text

    def test_render_sensitivity_empty(self):
        m = ModelMetrics()
        renderer = TerminalModelRenderer(use_color=False)
        text = renderer.render_sensitivity(m)
        assert "No data" in text


# ---------------------------------------------------------------------------
# render_from_assumptions
# ---------------------------------------------------------------------------

class TestRenderFromAssumptions:

    def test_default_assumptions(self):
        text = render_from_assumptions({})
        assert "Target Company" in text
        assert "MOIC" in text
        assert "IRR" in text

    def test_custom_assumptions(self):
        text = render_from_assumptions({
            "company_name": "Widget Inc",
            "ltm_revenue": 200.0,
            "ltm_ebitda": 40.0,
            "entry_multiple": 10.0,
            "exit_multiple": 10.0,
            "hold_period": 5,
        })
        assert "Widget Inc" in text
        assert "10.0x" in text

    def test_contains_all_sections(self):
        text = render_from_assumptions({"company_name": "Test"})
        assert "Sources & Uses" in text
        assert "Operating Model" in text
        assert "Debt Schedule" in text
        assert "Returns Analysis" in text
        assert "Sensitivity" in text


# ---------------------------------------------------------------------------
# extract_metrics with real workbook (mocked)
# ---------------------------------------------------------------------------

class TestExtractMetrics:

    def test_missing_assumptions_sheet_raises(self, tmp_path):
        """If no Assumptions sheet, raise ValueError."""
        try:
            from openpyxl import Workbook
        except ImportError:
            pytest.skip("openpyxl not installed")

        wb = Workbook()
        wb.active.title = "NotAssumptions"
        wb_path = tmp_path / "test.xlsx"
        wb.save(str(wb_path))

        with pytest.raises(ValueError, match="Assumptions"):
            extract_metrics(str(wb_path))

    def test_extract_from_generated_workbook(self, tmp_path):
        """Extract metrics from a workbook with an Assumptions sheet."""
        try:
            from openpyxl import Workbook
        except ImportError:
            pytest.skip("openpyxl not installed")

        wb = Workbook()
        ws = wb.active
        ws.title = "Assumptions"

        # Write the expected layout
        ws.cell(1, 1, "Test Corp - Model Assumptions")
        ws.cell(5, 2, 100.0)   # LTM Revenue
        ws.cell(6, 2, 20.0)    # LTM EBITDA
        ws.cell(7, 2, 8.0)     # Entry Multiple
        ws.cell(8, 2, 8.0)     # Exit Multiple
        ws.cell(9, 2, 5)       # Hold Period

        ws.cell(12, 2, 4.0)    # Senior Debt Multiple
        ws.cell(13, 2, 0.08)   # Senior Interest Rate
        ws.cell(14, 2, 0.01)   # Senior Amortization
        ws.cell(15, 2, 1.5)    # Sub Debt Multiple
        ws.cell(16, 2, 0.12)   # Sub Interest Rate

        ws.cell(19, 2, 0.015)  # Txn Fee %
        ws.cell(20, 2, 0.025)  # Fin Fee %
        ws.cell(22, 2, 0.05)   # CapEx %
        ws.cell(23, 2, 0.10)   # NWC %
        ws.cell(24, 2, 0.25)   # Tax Rate

        # Projection drivers (row 28 = rev growth, row 29 = EBITDA margin)
        for i, (g, m) in enumerate(zip(
            [0.08, 0.07, 0.06, 0.05, 0.05],
            [0.20, 0.21, 0.22, 0.22, 0.22],
        )):
            ws.cell(28, 3 + i, g)
            ws.cell(29, 3 + i, m)

        wb_path = tmp_path / "model.xlsx"
        wb.save(str(wb_path))

        m = extract_metrics(str(wb_path))

        assert m.ltm_revenue == 100.0
        assert m.ltm_ebitda == 20.0
        assert m.entry_multiple == 8.0
        assert m.hold_period == 5

        # Derived values should be computed
        assert m.entry_ev == pytest.approx(160.0)
        assert m.moic > 0
        assert m.irr > 0
        assert len(m.projected_revenue) == 5
        assert len(m.sensitivity_rows) > 0

    def test_full_render_from_workbook(self, tmp_path):
        """End-to-end: generate workbook, extract, render."""
        try:
            from openpyxl import Workbook
        except ImportError:
            pytest.skip("openpyxl not installed")

        wb = Workbook()
        ws = wb.active
        ws.title = "Assumptions"

        ws.cell(1, 1, "E2E Corp - Model Assumptions")
        ws.cell(5, 2, 150.0)
        ws.cell(6, 2, 30.0)
        ws.cell(7, 2, 10.0)
        ws.cell(8, 2, 9.0)
        ws.cell(9, 2, 5)
        ws.cell(12, 2, 4.0)
        ws.cell(13, 2, 0.07)
        ws.cell(14, 2, 0.02)
        ws.cell(15, 2, 1.5)
        ws.cell(16, 2, 0.10)
        ws.cell(19, 2, 0.02)
        ws.cell(20, 2, 0.03)
        ws.cell(22, 2, 0.04)
        ws.cell(23, 2, 0.08)
        ws.cell(24, 2, 0.25)

        for i, (g, m) in enumerate(zip(
            [0.10, 0.08, 0.06, 0.05, 0.05],
            [0.20, 0.21, 0.22, 0.23, 0.23],
        )):
            ws.cell(28, 3 + i, g)
            ws.cell(29, 3 + i, m)

        wb_path = tmp_path / "e2e.xlsx"
        wb.save(str(wb_path))

        renderer = TerminalModelRenderer(use_color=False)
        text = renderer.render_model(str(wb_path))

        assert "Sources & Uses" in text
        assert "Operating Model" in text
        assert "Returns Analysis" in text
        assert "Sensitivity" in text
        assert "MOIC" in text
