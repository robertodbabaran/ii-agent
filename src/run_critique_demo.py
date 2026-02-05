"""
Run the case study critique demo.
Execute from the src directory.
"""
import sys
sys.path.insert(0, '.')

import asyncio
from pathlib import Path
from datetime import datetime
from typing import Dict, Any

from ii_skills.shared.task_graph import (
    TaskGraph,
    Task,
    TaskGraphExecutor,
    LoggingCallbacks,
    ExecutionBudget,
)
from ii_skills.shared.event_telemetry import TelemetryLogger


# Case study data extracted from PDF
CASE_STUDY_DATA = {
    "title": "Optical Distributor Group Analysis",
    "date": "January 26, 2026",
    "company": "ODG (Optical Distributor Group)",
    "pages": {
        1: {
            "title": "Cover Page",
            "content": "VENTURE & GROWTH | OPTICAL DISTRIBUTOR GROUP ANALYSIS",
            "issues": ["Header says 'VENTURE & GROWTH' but this is an LBO/buyout transaction"]
        },
        2: {
            "title": "Company Overview",
            "content": "ODG is #2 U.S. Optical Distributor with industry-leading metrics",
            "issues": []
        },
        3: {
            "title": "Projections",
            "content": "8% Revenue CAGR with Margin Stability",
            "issues": [
                "Page contains 'NTD' placeholders - incomplete",
                "Shows 'Combo chart + table' placeholder instead of actual content",
                "Assessment of cash flows section not completed"
            ]
        },
        4: {
            "title": "Investment Recommendation",
            "content": "Qualified Buy recommendation",
            "issues": [
                "CRITICAL: Content is from WRONG case study - mentions QuickBooks, AI accounting, SaaS",
                "References 'Sasha' founder - not relevant to ODG optical distributor",
                "Mentions Series C at $856M valuation - ODG is a buyout, not VC",
                "Mentions 136% net dollar retention, 66% ARR growth - SaaS metrics, not distributor",
                "This entire page appears copy-pasted from a different analysis"
            ]
        },
        5: {"title": "Key Risks", "content": "Risk matrix with 5 key risks", "issues": []},
        6: {"title": "Valuation Benchmarking", "content": "Comps and precedents", "issues": []},
        7: {"title": "Valuation Analysis", "content": "$215M entry valuation", "issues": []},
        8: {"title": "Transaction Return Profile", "content": "Scenario analysis", "issues": []},
        9: {"title": "Due Diligence Approach", "content": "Structured DD plan", "issues": []}
    }
}


async def analyze_structure(context: Dict[str, Any]) -> Dict[str, Any]:
    """Analyze overall case structure."""
    data = context["case_data"]
    return {
        "total_pages": len(data["pages"]),
        "sections": [p["title"] for p in data["pages"].values()],
    }


async def analyze_content_quality(context: Dict[str, Any]) -> Dict[str, Any]:
    """Analyze content quality across pages."""
    data = context["case_data"]
    all_issues = []
    for page_num, page in data["pages"].items():
        for issue in page["issues"]:
            all_issues.append({
                "page": page_num,
                "section": page["title"],
                "issue": issue,
                "severity": "CRITICAL" if "CRITICAL" in issue else "WARNING"
            })
    return {
        "total_issues": len(all_issues),
        "critical_issues": sum(1 for i in all_issues if i["severity"] == "CRITICAL"),
        "issues": all_issues,
    }


async def analyze_consistency(context: Dict[str, Any]) -> Dict[str, Any]:
    """Check for consistency issues."""
    return {
        "inconsistencies": [
            {"type": "header_mismatch", "description": "Cover says 'Venture & Growth' but transaction is LBO"},
            {"type": "wrong_content", "page": 4, "description": "Investment Rec from different case"},
        ],
        "critical_count": 1,
    }


async def generate_powerpoint(context: Dict[str, Any]) -> str:
    """Generate the PowerPoint critique slide."""
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RgbColor

    output_path = context.get("output_path")

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(0.6))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = "Case Study Critique: ODG Analysis (v1)"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = RgbColor(0, 51, 102)

    # Grade
    grade_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.9), Inches(12), Inches(0.4))
    tf = grade_box.text_frame
    p = tf.paragraphs[0]
    p.text = "Overall Assessment: NEEDS REVISION"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = RgbColor(192, 0, 0)

    # Critical Issues (Left)
    issues_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(6), Inches(2.5))
    tf = issues_box.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Critical Issues"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = RgbColor(192, 0, 0)

    issues = [
        "Page 4: Investment recommendation is from a DIFFERENT case study (SaaS/AI accounting vs. Optical Distribution)",
        "Page 3: Projections page incomplete with 'NTD' placeholders",
        "Cover: Header says 'Venture & Growth' but this is an LBO transaction",
    ]
    for issue in issues:
        p = tf.add_paragraph()
        p.text = f"• {issue}"
        p.font.size = Pt(11)
        p.space_before = Pt(6)

    # Strengths (Right)
    strengths_box = slide.shapes.add_textbox(Inches(6.8), Inches(1.5), Inches(6), Inches(2.5))
    tf = strengths_box.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Strengths"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = RgbColor(0, 128, 0)

    strengths = [
        "Comprehensive risk analysis with probability/impact matrix",
        "Solid valuation benchmarking with comps and precedents",
        "Detailed return sensitivity analysis",
        "Thorough due diligence framework",
    ]
    for s in strengths:
        p = tf.add_paragraph()
        p.text = f"• {s}"
        p.font.size = Pt(11)
        p.space_before = Pt(6)

    # Recommendations (Bottom)
    rec_box = slide.shapes.add_textbox(Inches(0.5), Inches(4.2), Inches(12), Inches(2.5))
    tf = rec_box.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Recommended Actions Before Final Submission"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = RgbColor(0, 51, 102)

    recs = [
        "1. URGENT: Replace Page 4 content entirely - currently shows AI/SaaS case instead of ODG optical distribution",
        "2. Complete Page 3 projections - remove 'NTD' placeholders and add actual financial projections chart",
        "3. Update cover header from 'Venture & Growth' to 'Private Equity' or 'Buyout' to match deal type",
        "4. Add explicit investment recommendation for ODG with thesis points specific to optical distribution",
    ]
    for r in recs:
        p = tf.add_paragraph()
        p.text = r
        p.font.size = Pt(11)
        p.space_before = Pt(6)

    # Footer
    footer_box = slide.shapes.add_textbox(Inches(0.5), Inches(7.0), Inches(12), Inches(0.3))
    tf = footer_box.text_frame
    p = tf.paragraphs[0]
    p.text = f"Generated by Task Graph Demo | {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    p.font.size = Pt(9)
    p.font.color.rgb = RgbColor(128, 128, 128)

    prs.save(output_path)
    return output_path


async def run_critique():
    """Run the critique using task graph."""
    print("\n" + "="*60)
    print("TASK GRAPH DEMO: Case Study Critique")
    print("="*60)

    output_dir = Path("C:/Users/user/ii-agent/outputs/demos")
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / "ODG_Case_Critique.pptx"

    # Initialize telemetry
    telemetry = TelemetryLogger(
        output_dir=str(output_dir / "telemetry"),
        enable_console_logging=True,
    )

    # Create task graph
    graph = TaskGraph(name="CaseCritique")

    # Define phases
    graph.add_phase("analyze", parallel=True, max_parallel=3)
    graph.add_phase("output", parallel=False, depends_on=["analyze"])

    # Register handlers
    graph.register_handler("analyze_structure", analyze_structure)
    graph.register_handler("analyze_quality", analyze_content_quality)
    graph.register_handler("analyze_consistency", analyze_consistency)
    graph.register_handler("generate_pptx", generate_powerpoint)

    # Add analysis tasks (run in parallel)
    graph.add_task(Task(id="structure", name="Analyze Structure", phase="analyze", handler_name="analyze_structure"))
    graph.add_task(Task(id="quality", name="Analyze Quality", phase="analyze", handler_name="analyze_quality"))
    graph.add_task(Task(id="consistency", name="Check Consistency", phase="analyze", handler_name="analyze_consistency"))

    # Add output task (depends on analysis)
    graph.add_task(Task(
        id="pptx",
        name="Generate PowerPoint",
        phase="output",
        handler_name="generate_pptx",
        dependencies=["structure", "quality", "consistency"]
    ))

    # Create context
    context = {
        "case_data": CASE_STUDY_DATA,
        "output_path": str(output_path),
    }

    # Create budget
    budget = ExecutionBudget(max_total_tasks=10, max_parallel_tasks=3, task_timeout_ms=30000)

    # Execute with telemetry
    async with telemetry.run_context(run_type="case_critique", metadata={"case": "ODG v1"}) as run_ctx:
        executor = TaskGraphExecutor(graph=graph, budget=budget, callbacks=LoggingCallbacks(), context=context)
        results = await executor.run_all()

        run_ctx.increment_tasks(completed=len([r for r in results.values() if r.status.value == "completed"]))
        run_ctx.add_metric("output_file", str(output_path))

    print("\n" + "="*60)
    print("CRITIQUE RESULTS")
    print("="*60)
    print("\nGrade: NEEDS REVISION")
    print("\nTop Issues Found:")
    print("  1. Page 4 has WRONG content (SaaS case instead of ODG)")
    print("  2. Page 3 has 'NTD' placeholders - incomplete")
    print("  3. Cover says 'Venture & Growth' but this is an LBO")
    print(f"\nPowerPoint saved to: {output_path}")
    print("="*60)

    return str(output_path)


if __name__ == "__main__":
    result = asyncio.run(run_critique())
    print(f"\nDone! Output: {result}")
