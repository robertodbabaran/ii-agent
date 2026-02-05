"""
Excel Writer for Extracted Financial Data

Writes extracted financial data into professionally formatted Excel workbooks
following the conventions from EXCEL_CONVENTIONS.md.
"""

import logging
from pathlib import Path
from typing import Dict, List, Optional, Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

# Styles matching IB toolkit conventions
INPUT_FONT = Font(color="0000FF", bold=False)
CALC_FONT = Font(color="000000", bold=False)
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
SECTION_FONT = Font(bold=True, size=11)
TITLE_FONT = Font(bold=True, size=14)
SUBTOTAL_FONT = Font(bold=True)

HEADER_FILL = PatternFill(start_color="00365B", end_color="00365B", fill_type="solid")
SUBTOTAL_FILL = PatternFill(start_color="E8E8E8", end_color="E8E8E8", fill_type="solid")
INPUT_FILL = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")

THIN_BORDER = Border(bottom=Side(style="thin"))
DOUBLE_BORDER = Border(bottom=Side(style="double"))

# Canonical label display names
DISPLAY_NAMES = {
    # Income Statement
    "revenue": "Revenue",
    "cogs": "Cost of Goods Sold",
    "gross_profit": "Gross Profit",
    "sg_and_a": "SG&A",
    "rd_expense": "R&D Expense",
    "total_opex": "Total Operating Expenses",
    "ebitda": "EBITDA",
    "depreciation": "Depreciation & Amortization",
    "ebit": "EBIT / Operating Income",
    "interest_expense": "Interest Expense",
    "income_tax": "Income Tax",
    "net_income": "Net Income",
    # Balance Sheet
    "cash": "Cash & Equivalents",
    "accounts_receivable": "Accounts Receivable",
    "inventory": "Inventory",
    "total_current_assets": "Total Current Assets",
    "ppe": "PP&E",
    "goodwill": "Goodwill",
    "total_assets": "Total Assets",
    "accounts_payable": "Accounts Payable",
    "total_current_liabilities": "Total Current Liabilities",
    "long_term_debt": "Long-Term Debt",
    "total_liabilities": "Total Liabilities",
    "total_equity": "Total Equity",
    "total_liab_equity": "Total Liabilities & Equity",
    # Cash Flow
    "cfo": "Cash from Operations",
    "capex": "Capital Expenditures",
    "cfi": "Cash from Investing",
    "cff": "Cash from Financing",
    "fcf": "Free Cash Flow",
}

# Standard row ordering for each statement
INCOME_STATEMENT_ORDER = [
    "revenue", "cogs", "gross_profit", "sg_and_a", "rd_expense", "total_opex",
    "ebitda", "depreciation", "ebit", "interest_expense", "income_tax", "net_income",
]

BALANCE_SHEET_ORDER = [
    "cash", "accounts_receivable", "inventory", "total_current_assets",
    "ppe", "goodwill", "total_assets",
    "accounts_payable", "total_current_liabilities", "long_term_debt",
    "total_liabilities", "total_equity", "total_liab_equity",
]

CASH_FLOW_ORDER = ["cfo", "capex", "cfi", "cff", "fcf"]

# Items that should be styled as subtotals/totals
SUBTOTAL_ITEMS = {
    "gross_profit", "total_opex", "ebitda", "ebit", "net_income",
    "total_current_assets", "total_assets", "total_current_liabilities",
    "total_liabilities", "total_equity", "total_liab_equity",
    "cfo", "cfi", "cff", "fcf",
}

TOTAL_ITEMS = {"net_income", "total_assets", "total_liab_equity", "fcf"}


def write_statement_sheet(
    wb: Workbook,
    sheet_name: str,
    statement_title: str,
    data: Dict[str, Any],
    item_order: List[str],
    company_name: str = "",
) -> None:
    """Write a financial statement to an Excel sheet.

    Args:
        wb: Workbook to write to
        sheet_name: Name for the worksheet
        statement_title: Title displayed at top
        data: Extracted data with 'years' and 'line_items'
        item_order: Ordered list of canonical item names
        company_name: Company name for header
    """
    ws = wb.create_sheet(title=sheet_name)

    years = data.get("years", [])
    line_items = data.get("line_items", {})

    if not years or not line_items:
        ws["A1"] = f"No {statement_title} data extracted"
        ws["A1"].font = Font(italic=True, color="FF0000")
        return

    # Column setup
    label_col = 1
    first_data_col = 3  # Leave col B as spacer

    # -- Title Row --
    ws.cell(row=1, column=label_col, value=f"{company_name} — {statement_title}").font = TITLE_FONT

    # -- Year Headers (Row 3) --
    ws.cell(row=3, column=label_col, value="($ in millions)").font = Font(italic=True, size=9)
    for i, year in enumerate(years):
        cell = ws.cell(row=3, column=first_data_col + i, value=year)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")

    # -- Data Rows --
    current_row = 5

    for item_key in item_order:
        if item_key not in line_items:
            continue

        values = line_items[item_key]
        display_name = DISPLAY_NAMES.get(item_key, item_key.replace("_", " ").title())

        # Label cell
        label_cell = ws.cell(row=current_row, column=label_col, value=display_name)

        is_subtotal = item_key in SUBTOTAL_ITEMS
        is_total = item_key in TOTAL_ITEMS

        if is_total:
            label_cell.font = SUBTOTAL_FONT
            label_cell.border = DOUBLE_BORDER
        elif is_subtotal:
            label_cell.font = SUBTOTAL_FONT
            label_cell.fill = SUBTOTAL_FILL
            label_cell.border = THIN_BORDER

        # Value cells
        for i, value in enumerate(values):
            if i >= len(years):
                break
            cell = ws.cell(row=current_row, column=first_data_col + i)

            if value is not None:
                cell.value = value
                cell.number_format = '#,##0.0;(#,##0.0);"-"'
            else:
                cell.value = "-"

            if is_total:
                cell.font = SUBTOTAL_FONT
                cell.border = DOUBLE_BORDER
            elif is_subtotal:
                cell.font = SUBTOTAL_FONT
                cell.fill = SUBTOTAL_FILL
                cell.border = THIN_BORDER

            cell.alignment = Alignment(horizontal="right")

        current_row += 1

        # Add blank row after subtotals
        if is_subtotal:
            current_row += 1

    # -- Column widths --
    ws.column_dimensions[get_column_letter(label_col)].width = 35
    for i in range(len(years)):
        ws.column_dimensions[get_column_letter(first_data_col + i)].width = 14


def write_metrics_sheet(
    wb: Workbook,
    metrics: Dict[str, float],
    company_name: str = "",
) -> None:
    """Write extracted metrics to a summary sheet."""
    ws = wb.create_sheet(title="Key Metrics")

    ws.cell(row=1, column=1, value=f"{company_name} — Key Metrics").font = TITLE_FONT

    metric_display = {
        "revenue_growth": ("Revenue Growth", "0.0%"),
        "ebitda_margin": ("EBITDA Margin", "0.0%"),
        "gross_margin": ("Gross Margin", "0.0%"),
        "net_margin": ("Net Margin", "0.0%"),
        "capex_pct": ("CapEx % of Revenue", "0.0%"),
        "leverage": ("Net Leverage", "0.0x"),
        "debt_to_ebitda": ("Debt / EBITDA", "0.0x"),
    }

    row = 3
    for key, value in metrics.items():
        display_name, fmt = metric_display.get(key, (key.replace("_", " ").title(), "0.0"))

        ws.cell(row=row, column=1, value=display_name).font = Font(bold=True)

        cell = ws.cell(row=row, column=3, value=value)
        cell.number_format = fmt
        cell.font = INPUT_FONT
        cell.fill = INPUT_FILL
        cell.alignment = Alignment(horizontal="center")

        row += 1

    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["C"].width = 15


def write_extraction_to_excel(
    extracted_data: Dict[str, Any],
    output_path: str,
    company_name: str = "Company",
) -> str:
    """Write all extracted financial data to an Excel workbook.

    Args:
        extracted_data: Output from extract_all_financials()
        output_path: Path to save the Excel file
        company_name: Company name for headers

    Returns:
        Path to the saved Excel file
    """
    wb = Workbook()

    # Remove default sheet
    wb.remove(wb.active)

    # Write each financial statement
    if extracted_data.get("income_statement", {}).get("line_items"):
        write_statement_sheet(
            wb, "Income Statement", "Income Statement",
            extracted_data["income_statement"],
            INCOME_STATEMENT_ORDER,
            company_name,
        )

    if extracted_data.get("balance_sheet", {}).get("line_items"):
        write_statement_sheet(
            wb, "Balance Sheet", "Balance Sheet",
            extracted_data["balance_sheet"],
            BALANCE_SHEET_ORDER,
            company_name,
        )

    if extracted_data.get("cash_flow", {}).get("line_items"):
        write_statement_sheet(
            wb, "Cash Flow", "Cash Flow Statement",
            extracted_data["cash_flow"],
            CASH_FLOW_ORDER,
            company_name,
        )

    # Write metrics
    if extracted_data.get("metrics"):
        write_metrics_sheet(wb, extracted_data["metrics"], company_name)

    # Ensure at least one sheet exists
    if not wb.sheetnames:
        ws = wb.create_sheet("Summary")
        ws["A1"] = "No financial data could be extracted from the PDF"
        ws["A1"].font = Font(italic=True, color="FF0000")

    wb.save(output_path)
    return output_path
