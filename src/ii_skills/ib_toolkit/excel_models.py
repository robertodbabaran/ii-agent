#!/usr/bin/env python3
"""
Excel Financial Model Generator
Creates template Excel models for DCF, LBO, Comps, and 3-Statement models.
"""

import openpyxl
from openpyxl.styles import Font, Fill, PatternFill, Border, Side, Alignment, NamedStyle
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import FormulaRule
import os
from datetime import datetime

from config import OUTPUT_DIR, DEFAULT_WACC, DEFAULT_TERMINAL_GROWTH, DEFAULT_TAX_RATE, PROJECTION_YEARS

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_PATH = os.path.join(SCRIPT_DIR, OUTPUT_DIR)

# Styles
HEADER_FILL = PatternFill(start_color="1e3a5f", end_color="1e3a5f", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
INPUT_FILL = PatternFill(start_color="fff3cd", end_color="fff3cd", fill_type="solid")
INPUT_FONT = Font(color="000000", name="Calibri", size=11)
CALC_FONT = Font(color="000000", name="Calibri", size=11)
TITLE_FONT = Font(bold=True, color="1e3a5f", name="Calibri", size=14)
SECTION_FONT = Font(bold=True, color="1e3a5f", name="Calibri", size=11)
BORDER = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)


def style_header_row(ws, row, start_col, end_col):
    """Apply header styling to a row."""
    for col in range(start_col, end_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal='center')
        cell.border = BORDER


def style_input_cell(cell):
    """Style a cell as an input (yellow background)."""
    cell.fill = INPUT_FILL
    cell.font = INPUT_FONT
    cell.border = BORDER


def create_dcf_model(company_name: str = "Company") -> str:
    """Create a DCF valuation model template."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "DCF Model"

    # Title
    ws['A1'] = f"DCF Valuation Model - {company_name}"
    ws['A1'].font = TITLE_FONT
    ws.merge_cells('A1:H1')

    # Assumptions Section
    ws['A3'] = "KEY ASSUMPTIONS"
    ws['A3'].font = SECTION_FONT

    assumptions = [
        ("Tax Rate", DEFAULT_TAX_RATE, "E4"),
        ("WACC", DEFAULT_WACC, "E5"),
        ("Terminal Growth Rate", DEFAULT_TERMINAL_GROWTH, "E6"),
        ("Projection Years", PROJECTION_YEARS, "E7"),
    ]

    for i, (label, value, cell) in enumerate(assumptions, start=4):
        ws[f'D{i}'] = label
        ws[cell] = value
        style_input_cell(ws[cell])
        if "Rate" in label or "WACC" in label:
            ws[cell].number_format = '0.0%'

    # Historical and Projected Financials
    ws['A9'] = "INCOME STATEMENT PROJECTIONS"
    ws['A9'].font = SECTION_FONT

    # Years header
    base_year = datetime.now().year - 1
    years = [""] + [f"FY{base_year + i}" for i in range(PROJECTION_YEARS + 1)]

    for col, year in enumerate(years, start=1):
        ws.cell(row=10, column=col, value=year)
    style_header_row(ws, 10, 1, len(years))

    # Financial line items
    line_items = [
        "Revenue",
        "Revenue Growth %",
        "",
        "EBITDA",
        "EBITDA Margin %",
        "",
        "Depreciation & Amortization",
        "EBIT",
        "Taxes",
        "NOPAT (Net Operating Profit After Tax)",
        "",
        "Plus: D&A",
        "Less: CapEx",
        "Less: Change in NWC",
        "Unlevered Free Cash Flow",
    ]

    for i, item in enumerate(line_items, start=11):
        ws.cell(row=i, column=1, value=item)
        if item and "%" not in item:
            ws.cell(row=i, column=1).font = Font(bold=True) if item in ["Unlevered Free Cash Flow", "NOPAT (Net Operating Profit After Tax)"] else CALC_FONT

        # Add input cells for base year
        if item in ["Revenue", "EBITDA", "Depreciation & Amortization", "CapEx", "Change in NWC"]:
            cell = ws.cell(row=i, column=2)
            style_input_cell(cell)

    # DCF Calculation Section
    dcf_start = 11 + len(line_items) + 2
    ws.cell(row=dcf_start, column=1, value="DCF VALUATION")
    ws.cell(row=dcf_start, column=1).font = SECTION_FONT

    dcf_items = [
        ("Sum of PV of FCFs", "=SUM(calculated)"),
        ("Terminal Value", "=FCF*(1+g)/(WACC-g)"),
        ("PV of Terminal Value", "=TV/(1+WACC)^n"),
        ("Enterprise Value", "=Sum of PV + PV of TV"),
        ("", ""),
        ("Less: Total Debt", "Input"),
        ("Plus: Cash", "Input"),
        ("Equity Value", "=EV - Debt + Cash"),
        ("", ""),
        ("Shares Outstanding (M)", "Input"),
        ("Implied Share Price", "=Equity Value / Shares"),
    ]

    for i, (item, note) in enumerate(dcf_items, start=dcf_start + 1):
        ws.cell(row=i, column=1, value=item)
        ws.cell(row=i, column=2, value=note if "Input" in note else "")
        if "Input" in note:
            style_input_cell(ws.cell(row=i, column=2))
        if item in ["Enterprise Value", "Equity Value", "Implied Share Price"]:
            ws.cell(row=i, column=1).font = Font(bold=True)

    # Sensitivity Analysis Section
    sens_start = dcf_start + len(dcf_items) + 3
    ws.cell(row=sens_start, column=1, value="SENSITIVITY ANALYSIS")
    ws.cell(row=sens_start, column=1).font = SECTION_FONT

    ws.cell(row=sens_start + 1, column=1, value="Share Price Sensitivity to WACC and Terminal Growth")

    # Set column widths
    ws.column_dimensions['A'].width = 35
    for col in range(2, 10):
        ws.column_dimensions[get_column_letter(col)].width = 12

    # Save
    if not os.path.exists(OUTPUT_PATH):
        os.makedirs(OUTPUT_PATH)

    filename = f"DCF_Model_{company_name.replace(' ', '_')}_{datetime.now().strftime('%Y%m%d')}.xlsx"
    filepath = os.path.join(OUTPUT_PATH, filename)
    wb.save(filepath)
    print(f"DCF Model saved: {filepath}")
    return filepath


def create_lbo_model(company_name: str = "Target") -> str:
    """Create an LBO model template."""
    wb = openpyxl.Workbook()

    # Sources & Uses Tab
    ws_su = wb.active
    ws_su.title = "Sources & Uses"

    ws_su['A1'] = f"LBO Model - {company_name}"
    ws_su['A1'].font = TITLE_FONT

    ws_su['A3'] = "TRANSACTION ASSUMPTIONS"
    ws_su['A3'].font = SECTION_FONT

    txn_assumptions = [
        ("Purchase Price (Enterprise Value)", "", True),
        ("Equity Purchase Price", "", False),
        ("", "", False),
        ("LTM EBITDA", "", True),
        ("Entry EV/EBITDA Multiple", "", True),
        ("", "", False),
        ("Transaction Fees (% of EV)", "2.0%", True),
        ("Financing Fees (% of Debt)", "2.5%", True),
    ]

    for i, (label, default, is_input) in enumerate(txn_assumptions, start=4):
        ws_su.cell(row=i, column=1, value=label)
        cell = ws_su.cell(row=i, column=2, value=default)
        if is_input:
            style_input_cell(cell)

    # Sources
    ws_su['A14'] = "SOURCES OF FUNDS"
    ws_su['A14'].font = SECTION_FONT
    style_header_row(ws_su, 15, 1, 3)
    ws_su['A15'] = "Source"
    ws_su['B15'] = "Amount"
    ws_su['C15'] = "% of Total"

    sources = [
        "Senior Secured Term Loan",
        "Senior Unsecured Notes",
        "Subordinated Debt / Mezzanine",
        "Sponsor Equity",
        "Management Rollover",
        "Total Sources"
    ]
    for i, source in enumerate(sources, start=16):
        ws_su.cell(row=i, column=1, value=source)
        if source != "Total Sources":
            style_input_cell(ws_su.cell(row=i, column=2))

    # Uses
    ws_su['A24'] = "USES OF FUNDS"
    ws_su['A24'].font = SECTION_FONT
    style_header_row(ws_su, 25, 1, 3)
    ws_su['A25'] = "Use"
    ws_su['B25'] = "Amount"
    ws_su['C25'] = "% of Total"

    uses = [
        "Purchase Equity",
        "Refinance Existing Debt",
        "Transaction Fees",
        "Financing Fees",
        "Cash to Balance Sheet",
        "Total Uses"
    ]
    for i, use in enumerate(uses, start=26):
        ws_su.cell(row=i, column=1, value=use)

    # Operating Model Tab
    ws_op = wb.create_sheet("Operating Model")
    ws_op['A1'] = "OPERATING MODEL PROJECTIONS"
    ws_op['A1'].font = TITLE_FONT

    # Years
    years = [""] + [f"Year {i}" for i in range(6)]
    for col, year in enumerate(years, start=1):
        ws_op.cell(row=3, column=col, value=year)
    style_header_row(ws_op, 3, 1, len(years))

    op_items = [
        "Revenue",
        "Revenue Growth %",
        "",
        "EBITDA",
        "EBITDA Margin %",
        "",
        "Depreciation",
        "Amortization",
        "EBIT",
        "",
        "Interest Expense",
        "EBT",
        "Taxes",
        "Net Income",
        "",
        "EBITDA",
        "Less: CapEx",
        "Less: Change in NWC",
        "Less: Cash Interest",
        "Less: Cash Taxes",
        "Free Cash Flow for Debt Paydown",
    ]

    for i, item in enumerate(op_items, start=4):
        ws_op.cell(row=i, column=1, value=item)

    # Debt Schedule Tab
    ws_debt = wb.create_sheet("Debt Schedule")
    ws_debt['A1'] = "DEBT SCHEDULE"
    ws_debt['A1'].font = TITLE_FONT

    # Returns Tab
    ws_ret = wb.create_sheet("Returns Analysis")
    ws_ret['A1'] = "LBO RETURNS ANALYSIS"
    ws_ret['A1'].font = TITLE_FONT

    ws_ret['A3'] = "EXIT ASSUMPTIONS"
    ws_ret['A3'].font = SECTION_FONT

    exit_items = [
        ("Exit Year", 5, True),
        ("Exit EBITDA Multiple", "", True),
        ("Exit Enterprise Value", "", False),
        ("Less: Net Debt at Exit", "", False),
        ("Exit Equity Value", "", False),
    ]

    for i, (label, default, is_input) in enumerate(exit_items, start=4):
        ws_ret.cell(row=i, column=1, value=label)
        cell = ws_ret.cell(row=i, column=2, value=default)
        if is_input:
            style_input_cell(cell)

    ws_ret['A11'] = "RETURNS SUMMARY"
    ws_ret['A11'].font = SECTION_FONT

    returns_items = [
        "Initial Equity Investment",
        "Exit Equity Value",
        "",
        "Total Return (Multiple of Money)",
        "IRR",
    ]
    for i, item in enumerate(returns_items, start=12):
        ws_ret.cell(row=i, column=1, value=item)

    # Column widths
    for ws in [ws_su, ws_op, ws_debt, ws_ret]:
        ws.column_dimensions['A'].width = 35
        for col in range(2, 10):
            ws.column_dimensions[get_column_letter(col)].width = 12

    # Save
    filename = f"LBO_Model_{company_name.replace(' ', '_')}_{datetime.now().strftime('%Y%m%d')}.xlsx"
    filepath = os.path.join(OUTPUT_PATH, filename)
    wb.save(filepath)
    print(f"LBO Model saved: {filepath}")
    return filepath


def create_comps_model(company_name: str = "Company") -> str:
    """Create a trading/transaction comps model template."""
    wb = openpyxl.Workbook()

    # Trading Comps Tab
    ws_trading = wb.active
    ws_trading.title = "Trading Comps"

    ws_trading['A1'] = f"Comparable Companies Analysis - {company_name}"
    ws_trading['A1'].font = TITLE_FONT

    # Headers
    headers = [
        "Company", "Ticker", "Stock Price", "Shares Out", "Market Cap",
        "Net Debt", "Enterprise Value", "LTM Revenue", "LTM EBITDA",
        "EV/Revenue", "EV/EBITDA", "P/E"
    ]

    for col, header in enumerate(headers, start=1):
        ws_trading.cell(row=3, column=col, value=header)
    style_header_row(ws_trading, 3, 1, len(headers))

    # Sample rows
    for row in range(4, 14):
        for col in range(1, len(headers) + 1):
            style_input_cell(ws_trading.cell(row=row, column=col))

    # Summary stats
    ws_trading['A15'] = "Summary Statistics"
    ws_trading['A15'].font = SECTION_FONT

    stats = ["Mean", "Median", "High", "Low"]
    for i, stat in enumerate(stats, start=16):
        ws_trading.cell(row=i, column=1, value=stat)

    # Transaction Comps Tab
    ws_txn = wb.create_sheet("Transaction Comps")

    ws_txn['A1'] = "Precedent Transactions Analysis"
    ws_txn['A1'].font = TITLE_FONT

    txn_headers = [
        "Date", "Target", "Acquirer", "Transaction Value",
        "LTM Revenue", "LTM EBITDA", "EV/Revenue", "EV/EBITDA", "Premium Paid"
    ]

    for col, header in enumerate(txn_headers, start=1):
        ws_txn.cell(row=3, column=col, value=header)
    style_header_row(ws_txn, 3, 1, len(txn_headers))

    for row in range(4, 14):
        for col in range(1, len(txn_headers) + 1):
            style_input_cell(ws_txn.cell(row=row, column=col))

    # Valuation Summary Tab
    ws_val = wb.create_sheet("Valuation Summary")

    ws_val['A1'] = "Valuation Summary - Football Field"
    ws_val['A1'].font = TITLE_FONT

    val_methods = [
        ("52-Week Trading Range", "", "", ""),
        ("Trading Comps - EV/Revenue", "", "", ""),
        ("Trading Comps - EV/EBITDA", "", "", ""),
        ("Transaction Comps - EV/Revenue", "", "", ""),
        ("Transaction Comps - EV/EBITDA", "", "", ""),
        ("DCF Analysis", "", "", ""),
    ]

    ws_val['A3'] = "Methodology"
    ws_val['B3'] = "Low"
    ws_val['C3'] = "Mid"
    ws_val['D3'] = "High"
    style_header_row(ws_val, 3, 1, 4)

    for i, (method, low, mid, high) in enumerate(val_methods, start=4):
        ws_val.cell(row=i, column=1, value=method)
        style_input_cell(ws_val.cell(row=i, column=2))
        style_input_cell(ws_val.cell(row=i, column=3))
        style_input_cell(ws_val.cell(row=i, column=4))

    # Column widths
    for ws in [ws_trading, ws_txn, ws_val]:
        ws.column_dimensions['A'].width = 25
        for col in range(2, 15):
            ws.column_dimensions[get_column_letter(col)].width = 14

    # Save
    filename = f"Comps_Model_{company_name.replace(' ', '_')}_{datetime.now().strftime('%Y%m%d')}.xlsx"
    filepath = os.path.join(OUTPUT_PATH, filename)
    wb.save(filepath)
    print(f"Comps Model saved: {filepath}")
    return filepath


def create_three_statement_model(company_name: str = "Company") -> str:
    """Create a 3-statement integrated financial model."""
    wb = openpyxl.Workbook()

    # Assumptions Tab
    ws_assume = wb.active
    ws_assume.title = "Assumptions"

    ws_assume['A1'] = f"3-Statement Model - {company_name}"
    ws_assume['A1'].font = TITLE_FONT

    ws_assume['A3'] = "MODEL DRIVERS"
    ws_assume['A3'].font = SECTION_FONT

    drivers = [
        ("Revenue Growth Rate", "5.0%"),
        ("Gross Margin", "40.0%"),
        ("SG&A as % of Revenue", "20.0%"),
        ("R&D as % of Revenue", "10.0%"),
        ("D&A as % of Revenue", "5.0%"),
        ("Tax Rate", "25.0%"),
        ("CapEx as % of Revenue", "6.0%"),
        ("NWC as % of Revenue", "10.0%"),
    ]

    for i, (driver, default) in enumerate(drivers, start=4):
        ws_assume.cell(row=i, column=1, value=driver)
        cell = ws_assume.cell(row=i, column=2, value=default)
        style_input_cell(cell)

    # Income Statement Tab
    ws_is = wb.create_sheet("Income Statement")
    ws_is['A1'] = "INCOME STATEMENT"
    ws_is['A1'].font = TITLE_FONT

    years = [""] + ["Historical"] + [f"Year {i}" for i in range(1, 6)]
    for col, year in enumerate(years, start=1):
        ws_is.cell(row=3, column=col, value=year)
    style_header_row(ws_is, 3, 1, len(years))

    is_items = [
        "Revenue",
        "Cost of Goods Sold",
        "Gross Profit",
        "",
        "SG&A",
        "R&D",
        "Operating Expenses",
        "",
        "EBITDA",
        "Depreciation & Amortization",
        "EBIT (Operating Income)",
        "",
        "Interest Expense",
        "Interest Income",
        "EBT (Pre-Tax Income)",
        "",
        "Income Tax Expense",
        "Net Income",
    ]

    for i, item in enumerate(is_items, start=4):
        ws_is.cell(row=i, column=1, value=item)
        if item in ["Gross Profit", "EBITDA", "EBIT (Operating Income)", "EBT (Pre-Tax Income)", "Net Income"]:
            ws_is.cell(row=i, column=1).font = Font(bold=True)

    # Balance Sheet Tab
    ws_bs = wb.create_sheet("Balance Sheet")
    ws_bs['A1'] = "BALANCE SHEET"
    ws_bs['A1'].font = TITLE_FONT

    for col, year in enumerate(years, start=1):
        ws_bs.cell(row=3, column=col, value=year)
    style_header_row(ws_bs, 3, 1, len(years))

    bs_items = [
        "ASSETS",
        "Cash & Equivalents",
        "Accounts Receivable",
        "Inventory",
        "Prepaid Expenses",
        "Total Current Assets",
        "",
        "PP&E (Net)",
        "Intangibles",
        "Other Long-Term Assets",
        "Total Assets",
        "",
        "LIABILITIES",
        "Accounts Payable",
        "Accrued Expenses",
        "Current Portion of Debt",
        "Total Current Liabilities",
        "",
        "Long-Term Debt",
        "Other Long-Term Liabilities",
        "Total Liabilities",
        "",
        "SHAREHOLDERS' EQUITY",
        "Common Stock",
        "Retained Earnings",
        "Total Equity",
        "",
        "Total Liabilities & Equity",
        "",
        "Balance Check (should = 0)",
    ]

    for i, item in enumerate(bs_items, start=4):
        ws_bs.cell(row=i, column=1, value=item)
        if item in ["ASSETS", "LIABILITIES", "SHAREHOLDERS' EQUITY", "Total Current Assets",
                    "Total Assets", "Total Current Liabilities", "Total Liabilities",
                    "Total Equity", "Total Liabilities & Equity"]:
            ws_bs.cell(row=i, column=1).font = Font(bold=True)

    # Cash Flow Statement Tab
    ws_cf = wb.create_sheet("Cash Flow")
    ws_cf['A1'] = "CASH FLOW STATEMENT"
    ws_cf['A1'].font = TITLE_FONT

    for col, year in enumerate(years, start=1):
        ws_cf.cell(row=3, column=col, value=year)
    style_header_row(ws_cf, 3, 1, len(years))

    cf_items = [
        "OPERATING ACTIVITIES",
        "Net Income",
        "Depreciation & Amortization",
        "Change in Accounts Receivable",
        "Change in Inventory",
        "Change in Accounts Payable",
        "Change in Other Working Capital",
        "Cash from Operations",
        "",
        "INVESTING ACTIVITIES",
        "Capital Expenditures",
        "Acquisitions",
        "Cash from Investing",
        "",
        "FINANCING ACTIVITIES",
        "Debt Issuance / (Repayment)",
        "Dividends Paid",
        "Share Repurchases",
        "Cash from Financing",
        "",
        "Net Change in Cash",
        "Beginning Cash",
        "Ending Cash",
    ]

    for i, item in enumerate(cf_items, start=4):
        ws_cf.cell(row=i, column=1, value=item)
        if item in ["OPERATING ACTIVITIES", "INVESTING ACTIVITIES", "FINANCING ACTIVITIES",
                    "Cash from Operations", "Cash from Investing", "Cash from Financing",
                    "Net Change in Cash", "Ending Cash"]:
            ws_cf.cell(row=i, column=1).font = Font(bold=True)

    # Column widths
    for ws in [ws_assume, ws_is, ws_bs, ws_cf]:
        ws.column_dimensions['A'].width = 30
        for col in range(2, 10):
            ws.column_dimensions[get_column_letter(col)].width = 14

    # Save
    filename = f"Three_Statement_Model_{company_name.replace(' ', '_')}_{datetime.now().strftime('%Y%m%d')}.xlsx"
    filepath = os.path.join(OUTPUT_PATH, filename)
    wb.save(filepath)
    print(f"3-Statement Model saved: {filepath}")
    return filepath


if __name__ == "__main__":
    print("=" * 50)
    print("Excel Financial Model Generator")
    print("=" * 50)
    print()

    print("Creating model templates...")
    create_dcf_model("Sample Company")
    create_lbo_model("Sample Target")
    create_comps_model("Sample Company")
    create_three_statement_model("Sample Company")

    print()
    print(f"All models saved to: {OUTPUT_PATH}")
