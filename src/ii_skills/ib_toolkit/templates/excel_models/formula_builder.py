#!/usr/bin/env python3
"""
Formula Builder - Cell reference and Excel formula utilities.

Provides helper methods for constructing Excel formulas programmatically,
used by ExcelModelGenerator to write real formulas instead of hardcoded values.
"""

from openpyxl.utils import get_column_letter


class FormulaBuilder:
    """Utility class for building Excel formula strings."""

    @staticmethod
    def col_letter(col: int) -> str:
        """Convert 1-based column number to letter. col_letter(2) → 'B'"""
        return get_column_letter(col)

    @staticmethod
    def ref(row: int, col: int, abs_row: bool = False, abs_col: bool = False) -> str:
        """Cell reference: ref(5, 2) → 'B5', ref(5, 2, True, True) → '$B$5'"""
        c = get_column_letter(col)
        c_str = f"${c}" if abs_col else c
        r_str = f"${row}" if abs_row else str(row)
        return f"{c_str}{r_str}"

    @staticmethod
    def sheet_ref(sheet: str, row: int, col: int, abs_row: bool = False, abs_col: bool = False) -> str:
        """Cross-sheet reference: sheet_ref('Operating Model', 10, 3) → \"'Operating Model'!C10\""""
        c = get_column_letter(col)
        c_str = f"${c}" if abs_col else c
        r_str = f"${row}" if abs_row else str(row)
        return f"'{sheet}'!{c_str}{r_str}"

    @staticmethod
    def sum_range(start_row: int, end_row: int, col: int) -> str:
        """SUM formula: sum_range(4, 6, 2) → '=SUM(B4:B6)'"""
        c = get_column_letter(col)
        return f"=SUM({c}{start_row}:{c}{end_row})"

    @staticmethod
    def sum_row_range(row: int, start_col: int, end_col: int) -> str:
        """SUM across columns: sum_row_range(18, 2, 6) → '=SUM(B18:F18)'"""
        sc = get_column_letter(start_col)
        ec = get_column_letter(end_col)
        return f"=SUM({sc}{row}:{ec}{row})"

    @staticmethod
    def avg_cells(row1: int, col1: int, row2: int, col2: int) -> str:
        """Average of two cells: avg_cells(5, 2, 8, 2) → '=(B5+B8)/2'"""
        r1 = f"{get_column_letter(col1)}{row1}"
        r2 = f"{get_column_letter(col2)}{row2}"
        return f"=({r1}+{r2})/2"

    @staticmethod
    def pct(numerator_row: int, denominator_row: int, col: int) -> str:
        """Division formula: pct(13, 10, 3) → '=C13/C10'"""
        c = get_column_letter(col)
        return f"={c}{numerator_row}/{c}{denominator_row}"

    @staticmethod
    def iferror(formula: str, fallback: str = "0") -> str:
        """Wrap formula in IFERROR: iferror('B5/B3', '0') → '=IFERROR(B5/B3,0)'
        If formula starts with '=', the '=' is stripped from the inner formula."""
        inner = formula.lstrip("=")
        return f"=IFERROR({inner},{fallback})"

    @staticmethod
    def max_zero(formula: str) -> str:
        """Wrap formula in MAX(0,...): max_zero('C10+C11') → '=MAX(0,C10+C11)'"""
        inner = formula.lstrip("=")
        return f"=MAX(0,{inner})"

    # ============================================================
    # FINANCIAL FUNCTIONS
    # ============================================================

    @staticmethod
    def irr(cashflow_range: str) -> str:
        """IRR formula: irr('B5:G5') → '=IRR(B5:G5)'"""
        return f"=IRR({cashflow_range})"

    @staticmethod
    def xnpv(rate: str, values: str, dates: str) -> str:
        """XNPV formula: xnpv('B3','C5:G5','C4:G4') → '=XNPV(B3,C5:G5,C4:G4)'"""
        return f"=XNPV({rate},{values},{dates})"

    @staticmethod
    def npv(rate: str, values: str) -> str:
        """NPV formula: npv('B3','C5:G5') → '=NPV(B3,C5:G5)'"""
        return f"=NPV({rate},{values})"

    # ============================================================
    # ARRAY / LOOKUP FUNCTIONS
    # ============================================================

    @staticmethod
    def sumproduct(range1: str, range2: str) -> str:
        """SUMPRODUCT: sumproduct('B5:B10','C5:C10') → '=SUMPRODUCT(B5:B10,C5:C10)'"""
        return f"=SUMPRODUCT({range1},{range2})"

    @staticmethod
    def index_match(result_range: str, lookup_range: str, val: str) -> str:
        """INDEX/MATCH: index_match('C5:C20','B5:B20','\"Revenue\"') → '=INDEX(C5:C20,MATCH(\"Revenue\",B5:B20,0))'"""
        return f"=INDEX({result_range},MATCH({val},{lookup_range},0))"

    @staticmethod
    def vlookup(val: str, table: str, col: int, exact: bool = True) -> str:
        """VLOOKUP: vlookup('B3','A5:D20',3) → '=VLOOKUP(B3,A5:D20,3,FALSE)'"""
        match_type = "FALSE" if exact else "TRUE"
        return f"=VLOOKUP({val},{table},{col},{match_type})"

    # ============================================================
    # CONDITIONAL / LOGIC FUNCTIONS
    # ============================================================

    @staticmethod
    def if_formula(condition: str, true_val: str, false_val: str) -> str:
        """IF formula: if_formula('B5>0.2','"Pass"','"Fail"') → '=IF(B5>0.2,"Pass","Fail")'"""
        return f"=IF({condition},{true_val},{false_val})"

    @staticmethod
    def nested_if(conditions_values: list, else_val: str) -> str:
        """Nested IF chain from list of (condition, value) tuples.
        nested_if([('B5>0.3','"High"'),('B5>0.15','"Med"')],'"Low"')
        → '=IF(B5>0.3,"High",IF(B5>0.15,"Med","Low"))'"""
        if not conditions_values:
            return f"={else_val}"
        cond, val = conditions_values[0]
        if len(conditions_values) == 1:
            return f"=IF({cond},{val},{else_val})"
        # Build from inside out
        inner = else_val
        for cond, val in reversed(conditions_values):
            inner = f"IF({cond},{val},{inner})"
        return f"={inner}"

    # ============================================================
    # STATISTICAL / MATH FUNCTIONS
    # ============================================================

    @staticmethod
    def min_cells(*refs: str) -> str:
        """MIN of specific cells: min_cells('B5','C5','D5') → '=MIN(B5,C5,D5)'"""
        return f"=MIN({','.join(refs)})"

    @staticmethod
    def max_cells(*refs: str) -> str:
        """MAX of specific cells: max_cells('B5','C5','D5') → '=MAX(B5,C5,D5)'"""
        return f"=MAX({','.join(refs)})"

    @staticmethod
    def abs_val(ref: str) -> str:
        """ABS formula: abs_val('B5') → '=ABS(B5)'"""
        return f"=ABS({ref})"

    @staticmethod
    def round_val(ref: str, digits: int = 2) -> str:
        """ROUND formula: round_val('B5', 2) → '=ROUND(B5,2)'"""
        return f"=ROUND({ref},{digits})"

    @staticmethod
    def average_range(start_row: int, end_row: int, col: int) -> str:
        """AVERAGE formula: average_range(5, 10, 2) → '=AVERAGE(B5:B10)'"""
        c = get_column_letter(col)
        return f"=AVERAGE({c}{start_row}:{c}{end_row})"

    @staticmethod
    def median_range(range_str: str) -> str:
        """MEDIAN formula: median_range('B5:B20') → '=MEDIAN(B5:B20)'"""
        return f"=MEDIAN({range_str})"

    @staticmethod
    def percentile_range(range_str: str, pct: float) -> str:
        """PERCENTILE.INC: percentile_range('B5:B20', 0.25) → '=PERCENTILE.INC(B5:B20,0.25)'"""
        return f"=PERCENTILE.INC({range_str},{pct})"

    # ============================================================
    # RANGE REFERENCE HELPERS
    # ============================================================

    @staticmethod
    def range_ref(start_row: int, start_col: int, end_row: int, end_col: int) -> str:
        """Range reference: range_ref(5, 2, 10, 2) → 'B5:B10'"""
        sc = get_column_letter(start_col)
        ec = get_column_letter(end_col)
        return f"{sc}{start_row}:{ec}{end_row}"

    @staticmethod
    def col_range(start_row: int, end_row: int, col: int) -> str:
        """Column range: col_range(5, 10, 2) → 'B5:B10'"""
        c = get_column_letter(col)
        return f"{c}{start_row}:{c}{end_row}"
