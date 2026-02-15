#!/usr/bin/env python3
"""
Base module for Excel Model Generator.

Contains style constants, enums, registries, ModelAssumptions dataclass,
and ExcelModelGeneratorBase class with infrastructure methods.
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment, NamedStyle
from openpyxl.utils import get_column_letter
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.formatting.rule import DataBarRule, CellIsRule, ColorScaleRule, FormulaRule
from datetime import datetime
from .formula_builder import FormulaBuilder as FB
from dataclasses import dataclass, field
from typing import List, Dict, Optional, Any, Tuple
from enum import Enum
import os

# ============================================================
# STYLE CONSTANTS (Goldman Sachs / Institutional Standards)
# ============================================================

# Colors
INPUT_FILL = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")  # Light yellow
CALC_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")   # White
HEADER_FILL = PatternFill(start_color="00365B", end_color="00365B", fill_type="solid") # Navy
SUBTOTAL_FILL = PatternFill(start_color="E8E8E8", end_color="E8E8E8", fill_type="solid")
TOTAL_FILL = PatternFill(start_color="D0D0D0", end_color="D0D0D0", fill_type="solid")

# Fonts
INPUT_FONT = Font(color="0000FF", bold=False)      # Blue for inputs
CALC_FONT = Font(color="000000", bold=False)       # Black for formulas
HEADER_FONT = Font(color="FFFFFF", bold=True)      # White on navy
TITLE_FONT = Font(size=14, bold=True)
SECTION_FONT = Font(size=11, bold=True)

# Borders
THIN_BORDER = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)
BOTTOM_BORDER = Border(bottom=Side(style='thin'))
DOUBLE_BORDER = Border(bottom=Side(style='double'))


class ModelDepth(Enum):
    """Model complexity levels based on time available."""
    QUICK = "quick"           # 24-48 hours - essentials only
    STANDARD = "standard"     # 3-5 days - detailed analysis
    COMPREHENSIVE = "comprehensive"  # 7+ days - full institutional


# ============================================================
# MODEL MODULE REGISTRY
# ============================================================

MODEL_MODULES = {
    # Core Financial Modules
    "revenue_build": {
        "name": "Revenue Build",
        "description": "Revenue by segment/product with growth assumptions",
        "quick_version": "Top-line with growth rate",
        "standard_version": "Segment breakdown with drivers",
        "comprehensive_version": "Unit economics, cohorts, customer-level",
        "function": "add_revenue_build"
    },
    "expense_build": {
        "name": "Expense / SG&A Build",
        "description": "Cost structure breakdown (COGS, SG&A components)",
        "quick_version": "% of revenue assumptions",
        "standard_version": "Line-item breakdown",
        "comprehensive_version": "Headcount model, vendor analysis",
        "function": "add_expense_build"
    },
    "operating_model": {
        "name": "Operating Model",
        "description": "Income statement through EBITDA with projections",
        "quick_version": "Revenue → EBITDA (5 lines)",
        "standard_version": "Full P&L with margins",
        "comprehensive_version": "Segment P&Ls, detailed bridge",
        "function": "add_operating_model"
    },

    # Debt & Capital Structure
    "debt_schedule": {
        "name": "Debt Schedule",
        "description": "Debt tranches, amortization, interest, cash sweep",
        "quick_version": "Single tranche, simple paydown",
        "standard_version": "Multi-tranche with mandatory amort",
        "comprehensive_version": "Cash sweep, revolver, PIK, covenants",
        "function": "add_debt_schedule"
    },
    "wacc_calculation": {
        "name": "WACC Calculation",
        "description": "Cost of capital with component breakdown",
        "quick_version": "Simple WACC assumption",
        "standard_version": "CAPM with beta, debt cost",
        "comprehensive_version": "Unlevering/relevering, size premium",
        "function": "add_wacc_calculation"
    },

    # Working Capital
    "working_capital": {
        "name": "Working Capital Analysis",
        "description": "NWC components, days calculations, cash conversion",
        "quick_version": "NWC as % of revenue",
        "standard_version": "AR/AP/Inventory days",
        "comprehensive_version": "Seasonal patterns, normalization",
        "function": "add_working_capital"
    },

    # Transaction / LBO
    "sources_uses": {
        "name": "Sources & Uses",
        "description": "Transaction funding structure",
        "quick_version": "Simple S&U table",
        "standard_version": "Detailed with fees breakdown",
        "comprehensive_version": "Multiple scenarios, rollover equity",
        "function": "add_sources_uses"
    },
    "returns_analysis": {
        "name": "Returns Analysis",
        "description": "MOIC, IRR, value creation attribution",
        "quick_version": "MOIC and IRR calculation",
        "standard_version": "Exit year sensitivity, value bridge",
        "comprehensive_version": "Full attribution, management vs. buyer",
        "function": "add_returns_analysis"
    },
    "sensitivity_tables": {
        "name": "Sensitivity Analysis",
        "description": "2-way sensitivity matrices",
        "quick_version": "Entry/exit multiple matrix",
        "standard_version": "Multiple matrices (growth, margin)",
        "comprehensive_version": "Monte Carlo, tornado chart data",
        "function": "add_sensitivity_tables"
    },

    # Valuation
    "dcf_model": {
        "name": "DCF Valuation",
        "description": "Discounted cash flow analysis",
        "quick_version": "Terminal value + 5-year DCF",
        "standard_version": "Full FCF build, sensitivity",
        "comprehensive_version": "Multiple terminal methods, scenarios",
        "function": "add_dcf_valuation"
    },
    "trading_comps": {
        "name": "Trading Comps",
        "description": "Public comparable companies analysis",
        "quick_version": "5 comps with key multiples",
        "standard_version": "10+ comps with full metrics",
        "comprehensive_version": "Regression analysis, adjustments",
        "function": "add_trading_comps"
    },
    "transaction_comps": {
        "name": "Transaction Comps",
        "description": "Precedent transactions analysis",
        "quick_version": "5 deals with multiples",
        "standard_version": "10+ deals with context",
        "comprehensive_version": "Control premium analysis",
        "function": "add_transaction_comps"
    },

    # Advanced Modules
    "three_statement": {
        "name": "3-Statement Model",
        "description": "Linked IS, BS, CF statements",
        "quick_version": "Not recommended for quick cases",
        "standard_version": "Simplified linked statements",
        "comprehensive_version": "Full circular model",
        "function": "add_three_statement"
    },
    "covenant_analysis": {
        "name": "Covenant Analysis",
        "description": "Debt covenant compliance tracking",
        "quick_version": "Simple leverage test",
        "standard_version": "Leverage + coverage tests",
        "comprehensive_version": "Full covenant package, headroom",
        "function": "add_covenant_analysis"
    },
    "synergy_model": {
        "name": "Synergy Model",
        "description": "Revenue and cost synergy quantification",
        "quick_version": "High-level synergy estimate",
        "standard_version": "Phased synergy build",
        "comprehensive_version": "Bottom-up with implementation costs",
        "function": "add_synergy_model"
    },

    # CFA-Derived Modules
    "reverse_dcf": {
        "name": "Reverse DCF",
        "description": "Market-implied growth rates from current share price",
        "quick_version": "Single implied growth rate",
        "standard_version": "2-way table (WACC × terminal growth)",
        "comprehensive_version": "Full implied vs forecast comparison",
        "function": "add_reverse_dcf"
    },
    "dupont_analysis": {
        "name": "DuPont Analysis",
        "description": "ROE decomposition (3-component and 5-component)",
        "quick_version": "3-component DuPont",
        "standard_version": "5-component extended DuPont",
        "comprehensive_version": "Trend analysis with driver attribution",
        "function": "add_dupont_analysis"
    },
    "roic_decomposition": {
        "name": "ROIC Decomposition",
        "description": "Return on invested capital vs WACC spread",
        "quick_version": "ROIC calculation",
        "standard_version": "ROIC vs WACC with economic profit",
        "comprehensive_version": "Full decomposition with IC turnover",
        "function": "add_roic_decomposition"
    },
    "ddm_valuation": {
        "name": "Dividend Discount Model",
        "description": "DDM valuation (two-stage and H-model)",
        "quick_version": "Gordon Growth Model",
        "standard_version": "Two-stage DDM",
        "comprehensive_version": "H-model with sensitivity table",
        "function": "add_ddm_valuation"
    },
    "earnings_quality": {
        "name": "Earnings Quality Scores",
        "description": "Beneish M-Score, Altman Z-Score, Piotroski F-Score",
        "quick_version": "Single score summary",
        "standard_version": "All three scores with components",
        "comprehensive_version": "Full dashboard with conditional formatting",
        "function": "add_earnings_quality"
    },

    # Canonical Data Repository
    "source_index": {
        "name": "Source Index",
        "description": "Canonical repository of all metrics flowing into PowerPoint slides",
        "quick_version": "Numbered metrics with source links",
        "standard_version": "Full metrics with categories and slide references",
        "comprehensive_version": "Complete audit trail with data vintage and validation",
        "function": "add_source_index"
    },

    # CFA Sprint 4 Modules
    "tornado_sensitivity": {
        "name": "Tornado Sensitivity",
        "description": "Single-variable sensitivity with IRR impact ranking",
        "quick_version": "5 key variables with data bars",
        "standard_version": "7 variables with base/low/high IRR",
        "comprehensive_version": "Custom variables with linked assumptions",
        "function": "add_tornado_sensitivity"
    },
    "football_field": {
        "name": "Football Field",
        "description": "Valuation range comparison across methodologies",
        "quick_version": "3-method range chart",
        "standard_version": "5-method with weighted valuation",
        "comprehensive_version": "Full methodology with implied metrics",
        "function": "add_football_field"
    },
    "sotp_valuation": {
        "name": "SOTP Valuation",
        "description": "Sum-of-the-parts by business segment",
        "quick_version": "2-3 segments with blended multiple",
        "standard_version": "4+ segments with EV bridge",
        "comprehensive_version": "Full segment analysis with sensitivity",
        "function": "add_sotp_valuation"
    },
    "enhanced_wacc": {
        "name": "Enhanced WACC",
        "description": "Multi-method beta and cost of debt analysis",
        "quick_version": "Single WACC estimate",
        "standard_version": "3 beta methods, 3 Kd methods",
        "comprehensive_version": "Full range with sensitivity",
        "function": "add_enhanced_wacc"
    },
    "geographic_terminal_growth": {
        "name": "Geographic Terminal Growth",
        "description": "GDP-weighted terminal growth rate by geography",
        "quick_version": "2-3 regions with weighted growth",
        "standard_version": "5 regions with reasonability check",
        "comprehensive_version": "Country-level with adjustments",
        "function": "add_geographic_terminal_growth"
    },
    "multi_stage_dcf": {
        "name": "Multi-Stage DCF",
        "description": "3-stage DCF with transition period",
        "quick_version": "2-stage (high growth + terminal)",
        "standard_version": "3-stage with linear decay",
        "comprehensive_version": "Full projection with EV bridge",
        "function": "add_multi_stage_dcf"
    },

    # Core Extensions (from _core_modules.py)
    "scenario_analysis": {
        "name": "Scenario Analysis",
        "description": "Bull/Bear/Base case comparison with probability-weighted returns",
        "quick_version": "3 scenarios with key metric comparison",
        "standard_version": "Full P&L per scenario with probability weighting",
        "comprehensive_version": "Thesis-linked scenarios with cascading impact chains",
        "function": "add_scenario_analysis"
    },
    "management_vs_buyer": {
        "name": "Management vs Buyer Case",
        "description": "Side-by-side comparison of management projections vs buyer underwriting",
        "quick_version": "Revenue and EBITDA comparison",
        "standard_version": "Full P&L comparison with variance analysis",
        "comprehensive_version": "Returns bridge showing impact of each variance",
        "function": "add_management_vs_buyer"
    },

    # Due Diligence Modules (from _dd_modules.py)
    "quality_of_earnings": {
        "name": "Quality of Earnings",
        "description": "EBITDA normalization with adjustments and run-rate analysis",
        "quick_version": "Top 5 EBITDA adjustments",
        "standard_version": "Full adjustment bridge with categories",
        "comprehensive_version": "Pro forma run-rate with sensitivity on key adjustments",
        "function": "add_quality_of_earnings"
    },
    "working_capital_normalization": {
        "name": "Working Capital Normalization",
        "description": "NWC components, days analysis, peg mechanism",
        "quick_version": "NWC peg calculation",
        "standard_version": "Component-level days analysis with seasonal patterns",
        "comprehensive_version": "Full normalization with outlier identification",
        "function": "add_working_capital_normalization"
    },
    "customer_revenue_quality": {
        "name": "Customer & Revenue Quality",
        "description": "Customer concentration, retention, unit economics",
        "quick_version": "Top 10 customer concentration",
        "standard_version": "Retention cohorts and LTV analysis",
        "comprehensive_version": "Full customer-level revenue decomposition",
        "function": "add_customer_revenue_quality"
    },
    "credit_debt_sizing": {
        "name": "Credit & Debt Sizing",
        "description": "Credit ratios, debt capacity, stress testing",
        "quick_version": "Key leverage and coverage ratios",
        "standard_version": "Debt capacity analysis with agency benchmarks",
        "comprehensive_version": "Full stress test with downside scenarios",
        "function": "add_credit_debt_sizing"
    },

    # Capital Structure Modules (from _capital_structure.py)
    "dividend_recap": {
        "name": "Dividend Recapitalization",
        "description": "Mid-hold dividend, capital return, returns impact",
        "quick_version": "Simple recap sizing",
        "standard_version": "Recap with leverage constraints and returns impact",
        "comprehensive_version": "Multi-timing analysis with covenant headroom",
        "function": "add_dividend_recap"
    },
    "refinancing_analysis": {
        "name": "Refinancing Analysis",
        "description": "Rate savings, maturity extension, cost-benefit",
        "quick_version": "Rate comparison and annual savings",
        "standard_version": "Full restructuring with prepayment penalties",
        "comprehensive_version": "NPV of refinancing with multiple scenarios",
        "function": "add_refinancing_analysis"
    },
    "cap_table_waterfall": {
        "name": "Cap Table Waterfall",
        "description": "Equity distribution waterfall with LP/GP splits",
        "quick_version": "Simple equity split at exit",
        "standard_version": "Full waterfall with preferred return and catch-up",
        "comprehensive_version": "Multi-class equity with MIP and co-invest",
        "function": "add_cap_table_waterfall"
    },
    "sponsor_economics": {
        "name": "Sponsor Economics",
        "description": "GP carry, management fees, fund-level returns",
        "quick_version": "Deal-level carry calculation",
        "standard_version": "Fund-level economics with fees",
        "comprehensive_version": "Full GP P&L with clawback analysis",
        "function": "add_sponsor_economics"
    },

    # Transaction Modules (from _transaction_modules.py)
    "addon_analysis": {
        "name": "Add-on / Bolt-on Analysis",
        "description": "Platform plus add-on, synergies, combined returns",
        "quick_version": "Accretion summary",
        "standard_version": "Full combined model with synergies",
        "comprehensive_version": "Multi add-on pipeline with phased synergies",
        "function": "add_addon_analysis"
    },
    "carveout_analysis": {
        "name": "Carve-out Analysis",
        "description": "Standalone costs, stranded costs, TSA requirements",
        "quick_version": "Key standalone adjustments",
        "standard_version": "Full standalone P&L with TSA costs",
        "comprehensive_version": "Multi-year separation with stranded cost phase-out",
        "function": "add_carveout_analysis"
    },
    "earnout_model": {
        "name": "Earnout / Contingent Consideration",
        "description": "Performance milestones, probability-weighted value",
        "quick_version": "Milestone summary with target/max",
        "standard_version": "Probability-weighted earnout valuation",
        "comprehensive_version": "Monte Carlo on milestone achievement",
        "function": "add_earnout_model"
    },
    "purchase_price_allocation": {
        "name": "Purchase Price Allocation",
        "description": "Asset valuation, goodwill, intangibles",
        "quick_version": "Simplified PPA with goodwill residual",
        "standard_version": "Intangibles identification with useful lives",
        "comprehensive_version": "Full fair value with deferred tax impact",
        "function": "add_purchase_price_allocation"
    },

    # Value Creation Modules (from _value_creation.py)
    "value_creation_bridge": {
        "name": "Value Creation Bridge",
        "description": "EBITDA growth, multiple expansion, deleveraging attribution",
        "quick_version": "3-component bridge (growth, multiple, debt)",
        "standard_version": "5-component bridge with margin expansion",
        "comprehensive_version": "Full attribution with organic vs. inorganic",
        "function": "add_value_creation_bridge"
    },
    "hundred_day_plan": {
        "name": "100-Day Plan",
        "description": "Post-close priorities, quick wins, milestones",
        "quick_version": "Top 5 priorities",
        "standard_version": "Phased plan with owners and KPIs",
        "comprehensive_version": "Full Gantt with dependencies and risk flags",
        "function": "add_hundred_day_plan"
    },
    "exit_readiness": {
        "name": "Exit Readiness Assessment",
        "description": "Exit options, timing, value maximization",
        "quick_version": "Exit option comparison",
        "standard_version": "Readiness scorecard with gap analysis",
        "comprehensive_version": "Full exit planning with buyer universe mapping",
        "function": "add_exit_readiness"
    },
    "management_incentive_plan": {
        "name": "Management Incentive Plan",
        "description": "MIP structure, vesting, payout scenarios",
        "quick_version": "MIP pool size and allocation",
        "standard_version": "Vesting schedule with exit payout scenarios",
        "comprehensive_version": "Full MIP modeling with ratchets and sweet equity",
        "function": "add_management_incentive_plan"
    },

    # Specialized Modules (from _specialized_modules.py)
    "rollup_model": {
        "name": "Rollup / Platform Build",
        "description": "Multi-acquisition strategy, combined metrics",
        "quick_version": "Acquisition pipeline summary",
        "standard_version": "Combined model with multiple-arbitrage tracking",
        "comprehensive_version": "Full pipeline with individual deal returns",
        "function": "add_rollup_model"
    },
    "tax_analysis": {
        "name": "Tax Analysis",
        "description": "Tax structure, NOLs, step-up, effective rate",
        "quick_version": "Effective tax rate bridge",
        "standard_version": "NOL utilization and step-up analysis",
        "comprehensive_version": "Full tax structuring with jurisdiction analysis",
        "function": "add_tax_analysis"
    },
    "control_premium_analysis": {
        "name": "Control Premium Analysis",
        "description": "Premium to unaffected price, historical premiums",
        "quick_version": "Premium to 30-day VWAP",
        "standard_version": "Historical premium analysis with precedents",
        "comprehensive_version": "Full premium decomposition with accretion/dilution",
        "function": "add_control_premium_analysis"
    },
}


# ============================================================
# TIME-BASED MODULE RECOMMENDATIONS
# ============================================================

CASE_FRAMEWORKS = {
    "24_hour": {
        "name": "24-Hour Sprint Case",
        "description": "Essential analysis only - focus on key decision drivers",
        "required_modules": [
            "sources_uses",
            "operating_model",  # Quick version
            "returns_analysis",
        ],
        "optional_modules": [
            "sensitivity_tables",  # Entry/exit only
        ],
        "depth": ModelDepth.QUICK,
        "tips": [
            "Use management projections directly if provided",
            "Simple debt structure (single tranche)",
            "Focus on MOIC/IRR, not detailed attribution",
            "One sensitivity table maximum",
        ]
    },
    "48_hour": {
        "name": "48-Hour Standard Case",
        "description": "Solid analysis with key supporting detail",
        "required_modules": [
            "sources_uses",
            "operating_model",
            "debt_schedule",
            "returns_analysis",
            "sensitivity_tables",
        ],
        "optional_modules": [
            "revenue_build",
            "working_capital",
            "dcf_model",
        ],
        "depth": ModelDepth.QUICK,
        "tips": [
            "Challenge 1-2 management assumptions",
            "Multi-tranche debt if relevant",
            "2-3 sensitivity matrices",
            "Simple value creation bridge",
        ]
    },
    "5_day": {
        "name": "5-Day Detailed Case",
        "description": "Comprehensive analysis with supporting detail",
        "required_modules": [
            "sources_uses",
            "revenue_build",
            "expense_build",
            "operating_model",
            "debt_schedule",
            "working_capital",
            "returns_analysis",
            "sensitivity_tables",
        ],
        "optional_modules": [
            "dcf_model",
            "trading_comps",
            "covenant_analysis",
            "wacc_calculation",
            "reverse_dcf",
            "dupont_analysis",
            "tornado_sensitivity",
            "football_field",
            "sotp_valuation",
            "earnings_quality",
        ],
        "depth": ModelDepth.STANDARD,
        "tips": [
            "Build revenue from segment drivers",
            "Detailed SG&A breakdown",
            "Cash sweep mechanics",
            "Multiple scenarios (base/upside/downside)",
            "Full value creation attribution",
        ]
    },
    "7_day_plus": {
        "name": "7+ Day Institutional Case",
        "description": "Full institutional-quality model",
        "required_modules": [
            "sources_uses",
            "revenue_build",
            "expense_build",
            "operating_model",
            "debt_schedule",
            "working_capital",
            "wacc_calculation",
            "returns_analysis",
            "sensitivity_tables",
            "dcf_model",
            "trading_comps",
            "transaction_comps",
            "covenant_analysis",
        ],
        "optional_modules": [
            "three_statement",
            "synergy_model",
            "reverse_dcf",
            "dupont_analysis",
            "roic_decomposition",
            "ddm_valuation",
            "earnings_quality",
            "tornado_sensitivity",
            "football_field",
            "sotp_valuation",
            "enhanced_wacc",
            "geographic_terminal_growth",
            "multi_stage_dcf",
            "scenario_analysis",
            "management_vs_buyer",
        ],
        "depth": ModelDepth.COMPREHENSIVE,
        "tips": [
            "Unit economics and cohort analysis",
            "Headcount-driven SG&A model",
            "Complex debt with revolver, PIK",
            "Management case vs. buyer case comparison",
            "Full covenant analysis with headroom",
            "Monte Carlo or scenario probability weighting",
        ]
    },
}


@dataclass
class ModelAssumptions:
    """Container for model assumptions."""
    # Company info
    company_name: str = "Target Company"
    fiscal_year_end: str = "December"

    # Historical
    ltm_revenue: float = 100.0
    ltm_ebitda: float = 20.0
    ltm_capex: float = 5.0

    # Growth assumptions
    revenue_growth: List[float] = field(default_factory=lambda: [0.08, 0.07, 0.06, 0.05, 0.05])
    ebitda_margin: List[float] = field(default_factory=lambda: [0.20, 0.21, 0.22, 0.22, 0.22])
    capex_pct_revenue: float = 0.05
    nwc_pct_revenue: float = 0.10

    # Transaction
    entry_multiple: float = 8.0
    exit_multiple: float = 8.0
    hold_period: int = 5

    # Debt structure
    senior_debt_multiple: float = 4.0
    senior_interest_rate: float = 0.08
    senior_amortization: float = 0.01  # 1% per year
    sub_debt_multiple: float = 1.5
    sub_interest_rate: float = 0.12

    # Fees
    transaction_fee_pct: float = 0.015  # 1.5% of EV
    financing_fee_pct: float = 0.025    # 2.5% of debt

    # WACC inputs
    risk_free_rate: float = 0.04
    equity_risk_premium: float = 0.06
    beta: float = 1.0
    size_premium: float = 0.02
    cost_of_debt: float = 0.08
    tax_rate: float = 0.25
    target_debt_equity: float = 0.50


class ExcelModelGeneratorBase:
    """
    Base class for the modular Excel model generator.
    Contains __init__, helper methods, build_for_timeframe, save.

    The full ExcelModelGenerator class inherits from this base plus all
    mixin classes, assembled in excel_modules.py.
    """

    def __init__(self, company_name: str = "Target Company",
                 depth: ModelDepth = ModelDepth.STANDARD,
                 assumptions: ModelAssumptions = None):
        """Initialize the Excel model generator."""
        self.wb = Workbook()
        self.company_name = company_name
        self.depth = depth
        self.assumptions = assumptions or ModelAssumptions(company_name=company_name)
        self.projection_years = self.assumptions.hold_period

        # Remove default sheet
        self.wb.remove(self.wb.active)

        # Track created sheets
        self.sheets_created = []

        # Track key cell locations across sheets for cross-references
        self.cell_map = {}  # e.g., {"assumptions": {"ltm_ebitda": "B6", ...}}

    def _format_header_row(self, ws, row: int, start_col: int, end_col: int):
        """Apply header formatting to a row."""
        for col in range(start_col, end_col + 1):
            cell = ws.cell(row=row, column=col)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal='center')

    def _format_input_cell(self, ws, row: int, col: int):
        """Format a cell as an input."""
        cell = ws.cell(row=row, column=col)
        cell.fill = INPUT_FILL
        cell.font = INPUT_FONT

    def _add_title(self, ws, title: str, row: int = 1, col: int = 1):
        """Add a section title."""
        cell = ws.cell(row=row, column=col, value=title)
        cell.font = TITLE_FONT

    def _add_section_header(self, ws, title: str, row: int, col: int = 1):
        """Add a section header."""
        cell = ws.cell(row=row, column=col, value=title)
        cell.font = SECTION_FONT

    def _year_headers(self, ws, start_row: int, start_col: int):
        """Add year column headers."""
        headers = ["LTM"] + [f"Year {i}" for i in range(1, self.projection_years + 1)]
        for i, header in enumerate(headers):
            cell = ws.cell(row=start_row, column=start_col + i, value=header)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal='center')

    # ============================================================
    # CONDITIONAL FORMATTING HELPERS
    # ============================================================

    def _add_rag_cells(self, ws, cell_range: str,
                       green_threshold, amber_threshold=None,
                       higher_is_better: bool = True):
        """Apply Red/Amber/Green conditional formatting to a cell range.

        Args:
            ws: Worksheet object
            cell_range: e.g. 'C12:G12'
            green_threshold: Value above (or below) which is green
            amber_threshold: Optional middle threshold for amber zone
            higher_is_better: True = green for high values, red for low
        """
        green = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
        red = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
        amber = PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid')

        if higher_is_better:
            ws.conditional_formatting.add(
                cell_range,
                CellIsRule(operator='greaterThanOrEqual',
                           formula=[str(green_threshold)], fill=green))
            if amber_threshold is not None:
                ws.conditional_formatting.add(
                    cell_range,
                    CellIsRule(operator='between',
                               formula=[str(amber_threshold), str(green_threshold)],
                               fill=amber))
                ws.conditional_formatting.add(
                    cell_range,
                    CellIsRule(operator='lessThan',
                               formula=[str(amber_threshold)], fill=red))
            else:
                ws.conditional_formatting.add(
                    cell_range,
                    CellIsRule(operator='lessThan',
                               formula=[str(green_threshold)], fill=red))
        else:
            # Lower is better (e.g., leverage ratio)
            ws.conditional_formatting.add(
                cell_range,
                CellIsRule(operator='lessThanOrEqual',
                           formula=[str(green_threshold)], fill=green))
            if amber_threshold is not None:
                ws.conditional_formatting.add(
                    cell_range,
                    CellIsRule(operator='between',
                               formula=[str(green_threshold), str(amber_threshold)],
                               fill=amber))
                ws.conditional_formatting.add(
                    cell_range,
                    CellIsRule(operator='greaterThan',
                               formula=[str(amber_threshold)], fill=red))
            else:
                ws.conditional_formatting.add(
                    cell_range,
                    CellIsRule(operator='greaterThan',
                               formula=[str(green_threshold)], fill=red))

    def _add_data_bars(self, ws, cell_range: str, color: str = '638EC6'):
        """Apply data bar conditional formatting to a cell range."""
        ws.conditional_formatting.add(
            cell_range,
            DataBarRule(start_type='min', end_type='max',
                        color=color, showValue=True))

    def _add_color_scale(self, ws, cell_range: str,
                          low_color: str = 'F8696B',
                          mid_color: str = 'FFEB84',
                          high_color: str = '63BE7B'):
        """Apply 3-color scale (red-yellow-green) to a cell range."""
        ws.conditional_formatting.add(
            cell_range,
            ColorScaleRule(start_type='min', start_color=low_color,
                           mid_type='percentile', mid_value=50, mid_color=mid_color,
                           end_type='max', end_color=high_color))

    # ============================================================
    # CASE FRAMEWORK GENERATION
    # ============================================================

    def build_for_timeframe(self, timeframe: str) -> 'ExcelModelGeneratorBase':
        """
        Build a model based on available time.

        Args:
            timeframe: One of "24_hour", "48_hour", "5_day", "7_day_plus"
        """
        if timeframe not in CASE_FRAMEWORKS:
            raise ValueError(f"Unknown timeframe: {timeframe}. Use: {list(CASE_FRAMEWORKS.keys())}")

        framework = CASE_FRAMEWORKS[timeframe]
        self.depth = framework["depth"]

        # Always add Assumptions sheet first (central input sheet for formulas)
        self.add_assumptions_sheet()

        # Add required modules
        for module in framework["required_modules"]:
            method_name = f"add_{module}"
            if hasattr(self, method_name):
                getattr(self, method_name)()

        return self

    # ============================================================
    # OUTPUT
    # ============================================================

    def save(self, filepath: str) -> str:
        """Save the workbook to file."""
        # Add cover sheet
        cover = self.wb.create_sheet("Cover", 0)
        cover.cell(row=2, column=2, value=self.company_name)
        cover.cell(row=2, column=2).font = Font(size=24, bold=True)
        cover.cell(row=4, column=2, value="Financial Model")
        cover.cell(row=4, column=2).font = Font(size=18)
        cover.cell(row=6, column=2, value=f"Generated: {datetime.now().strftime('%B %d, %Y')}")
        cover.cell(row=8, column=2, value=f"Depth: {self.depth.value.title()}")
        cover.cell(row=10, column=2, value="Sheets:")
        for i, sheet in enumerate(self.sheets_created):
            cover.cell(row=11 + i, column=2, value=f"  • {sheet}")

        cover.column_dimensions['B'].width = 40

        self.wb.save(filepath)
        print(f"Model saved: {filepath}")
        print(f"Sheets created: {len(self.sheets_created)}")
        return filepath

    def get_workbook(self) -> Workbook:
        """Return the raw workbook for further manipulation."""
        return self.wb


# ============================================================
# HELPER FUNCTIONS
# ============================================================

def list_model_modules() -> Dict:
    """Return the registry of available model modules."""
    return MODEL_MODULES


def list_case_frameworks() -> Dict:
    """Return the available case frameworks by time."""
    return CASE_FRAMEWORKS


def get_framework_recommendation(hours_available: int) -> Dict:
    """Get recommended framework based on hours available."""
    if hours_available <= 24:
        return CASE_FRAMEWORKS["24_hour"]
    elif hours_available <= 48:
        return CASE_FRAMEWORKS["48_hour"]
    elif hours_available <= 120:  # 5 days
        return CASE_FRAMEWORKS["5_day"]
    else:
        return CASE_FRAMEWORKS["7_day_plus"]
