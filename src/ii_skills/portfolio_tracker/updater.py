#!/usr/bin/env python3
"""
3-Mode Updater for Portfolio Tracker

Modes:
1. full_sync  — Refresh trades + prices + reference data + regenerate workbook
2. quick_update — Append latest prices to Market_Data in existing workbook
3. add_trade — Manual trade entry → recompute units/cash → regenerate
"""

import os
from datetime import datetime
from typing import Dict, List, Optional

from openpyxl import load_workbook

from .data_fetcher import fetch_current_price, fetch_exchange_rate
from .trade_log import Trade
from .workbook_generator import PortfolioWorkbookGenerator


class PortfolioUpdater:
    """
    Manages the 3 update modes for the portfolio tracker workbook.

    Usage:
        updater = PortfolioUpdater(output_dir="output/")
        updater.full_sync()
        updater.quick_update(workbook_path="output/Portfolio_Tracker.xlsx")
        updater.add_trade(ticker="AAPL", action="BUY", quantity=10, price=150.0)
    """

    def __init__(
        self,
        output_dir: str = "output",
        benchmark: str = "SPY",
        risk_free_rate: float = 0.045,
        target_allocation: Optional[Dict[str, float]] = None,
    ):
        self.output_dir = output_dir
        self.benchmark = benchmark
        self.risk_free_rate = risk_free_rate
        self.target_allocation = target_allocation or {}

    def _get_generator(self) -> PortfolioWorkbookGenerator:
        """Create a fresh generator instance."""
        return PortfolioWorkbookGenerator(
            output_dir=self.output_dir,
            benchmark=self.benchmark,
            risk_free_rate=self.risk_free_rate,
            target_allocation=self.target_allocation,
        )

    def _find_latest_workbook(self) -> Optional[str]:
        """Find the most recent workbook in output directory."""
        if not os.path.isdir(self.output_dir):
            return None
        files = [
            f for f in os.listdir(self.output_dir)
            if f.startswith("Portfolio_Tracker") and f.endswith(".xlsx")
        ]
        if not files:
            return None
        files.sort(reverse=True)
        return os.path.join(self.output_dir, files[0])

    # ────────────────────────────────────────────────────────────
    # Mode 1: Full Sync
    # ────────────────────────────────────────────────────────────

    def full_sync(self, output_name: Optional[str] = None,
                  period: str = "1y", **kwargs) -> Dict:
        """Full sync: regenerate entire workbook with fresh data."""
        generator = self._get_generator()
        output_path = generator.generate(output_name=output_name, period=period)
        return {
            "success": True,
            "mode": "full_sync",
            "output_path": output_path,
            "sheets": generator.sheet_names,
        }

    # ────────────────────────────────────────────────────────────
    # Mode 2: Quick Update
    # ────────────────────────────────────────────────────────────

    def quick_update(self, workbook_path: Optional[str] = None, **kwargs) -> Dict:
        """
        Quick update: append latest prices to Market_Data in existing workbook.

        Opens the workbook, finds Market_Data, appends a new row with
        today's prices from yfinance.
        """
        if workbook_path is None:
            workbook_path = self._find_latest_workbook()
        if not workbook_path or not os.path.exists(workbook_path):
            return {
                "success": False,
                "error": "No workbook found. Run generate_workbook first.",
            }

        try:
            wb = load_workbook(workbook_path)
        except Exception as e:
            return {"success": False, "error": f"Could not open workbook: {e}"}

        if "Market_Data" not in wb.sheetnames:
            return {"success": False, "error": "Market_Data sheet not found"}

        ws = wb["Market_Data"]

        # Read ticker columns from header row
        headers = []
        col = 1
        while True:
            val = ws.cell(row=1, column=col).value
            if val is None:
                break
            headers.append(val)
            col += 1

        if len(headers) < 2:
            return {"success": False, "error": "Market_Data headers not found"}

        # Find the next empty row
        next_row = ws.max_row + 1

        # Write today's date
        ws.cell(row=next_row, column=1, value=datetime.now())
        ws.cell(row=next_row, column=1).number_format = "YYYY-MM-DD"

        # Fetch prices for each ticker column
        updated_tickers = []
        for c_idx, ticker in enumerate(headers[1:], 2):  # skip Date column
            if ticker == "USDCAD":
                price = fetch_exchange_rate()
            else:
                price = fetch_current_price(ticker)

            if price is not None:
                ws.cell(row=next_row, column=c_idx, value=price)
                ws.cell(row=next_row, column=c_idx).number_format = "#,##0.00"
                updated_tickers.append(ticker)
            else:
                # Copy previous row's value as fallback
                prev_val = ws.cell(row=next_row - 1, column=c_idx).value
                ws.cell(row=next_row, column=c_idx, value=prev_val or 0)

        wb.save(workbook_path)
        return {
            "success": True,
            "mode": "quick_update",
            "workbook_path": workbook_path,
            "row_added": next_row,
            "tickers_updated": updated_tickers,
        }

    # ────────────────────────────────────────────────────────────
    # Mode 3: Add Trade
    # ────────────────────────────────────────────────────────────

    def add_trade(self, ticker: str, action: str, quantity: float,
                  price: float, currency: str = "CAD",
                  account: str = "", fees: float = 0.0,
                  date: Optional[str] = None, **kwargs) -> Dict:
        """
        Add a manual trade entry and regenerate workbook.

        After adding the trade to Trade_Log, the entire workbook is
        regenerated so Daily_Units, Daily_Cash, and all analytics update.
        """
        # Parse date
        if date:
            try:
                trade_date = datetime.strptime(date, "%Y-%m-%d")
            except ValueError:
                return {"success": False, "error": f"Invalid date format: {date}. Use YYYY-MM-DD."}
        else:
            trade_date = datetime.now()

        # Validate action
        action = action.upper()
        if action not in ("BUY", "SELL", "DIVIDEND"):
            return {"success": False, "error": f"Invalid action: {action}. Use BUY, SELL, or DIVIDEND."}

        trade = Trade(
            date=trade_date,
            ticker=ticker.upper(),
            action=action,
            quantity=float(quantity),
            price=float(price),
            currency=currency.upper(),
            fees=float(fees),
            account=account,
            notes=f"Manual entry via add_trade",
        )

        # Generate fresh workbook with the new trade
        generator = self._get_generator()
        generator._load_all_data("1y")

        # Inject the trade
        generator._trade_log.add_trade(trade)

        # Recompute daily units and cash
        dates = generator._price_history.dates
        generator._daily_units = generator._trade_log.compute_daily_units(dates)
        generator._daily_cash = generator._trade_log.compute_daily_cash(dates)

        # If the ticker is new, we need its price data
        if trade.ticker not in generator._price_history.prices:
            from .data_fetcher import fetch_price_history, fetch_reference_data
            new_prices = fetch_price_history([trade.ticker], "1y")
            if trade.ticker in new_prices.prices:
                generator._price_history.prices[trade.ticker] = new_prices.prices[trade.ticker]
            new_ref = fetch_reference_data([trade.ticker])
            generator._ref_data.update(new_ref)
            generator._tickers = generator._trade_log.get_tickers()

        # Regenerate
        output_path = generator.generate()

        return {
            "success": True,
            "mode": "add_trade",
            "trade": {
                "ticker": trade.ticker,
                "action": trade.action,
                "quantity": trade.quantity,
                "price": trade.price,
                "date": trade.date.strftime("%Y-%m-%d"),
            },
            "output_path": output_path,
        }

    # ────────────────────────────────────────────────────────────
    # Update Theses
    # ────────────────────────────────────────────────────────────

    def update_theses(self, ticker: str, thesis: str = "",
                      catalyst: str = "", edge: str = "",
                      workbook_path: Optional[str] = None, **kwargs) -> Dict:
        """Update investment thesis text for a ticker in the existing workbook."""
        if workbook_path is None:
            workbook_path = self._find_latest_workbook()
        if not workbook_path or not os.path.exists(workbook_path):
            return {"success": False, "error": "No workbook found. Run generate_workbook first."}

        try:
            wb = load_workbook(workbook_path)
        except Exception as e:
            return {"success": False, "error": f"Could not open workbook: {e}"}

        if "Investment_Theses" not in wb.sheetnames:
            return {"success": False, "error": "Investment_Theses sheet not found"}

        ws = wb["Investment_Theses"]

        # Find the ticker row
        target_row = None
        for row in range(2, ws.max_row + 1):
            if ws.cell(row=row, column=1).value == ticker.upper():
                target_row = row
                break

        if target_row is None:
            return {"success": False, "error": f"Ticker {ticker} not found in Investment_Theses"}

        if thesis:
            ws.cell(row=target_row, column=3, value=thesis)
        if catalyst:
            ws.cell(row=target_row, column=4, value=catalyst)
        if edge:
            ws.cell(row=target_row, column=5, value=edge)

        wb.save(workbook_path)
        return {
            "success": True,
            "ticker": ticker.upper(),
            "workbook_path": workbook_path,
            "updated": {
                "thesis": bool(thesis),
                "catalyst": bool(catalyst),
                "edge": bool(edge),
            },
        }
