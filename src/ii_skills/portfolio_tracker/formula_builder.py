#!/usr/bin/env python3
"""
Extended Formula Builder for Portfolio Tracker

Extends the IB Toolkit FormulaBuilder with portfolio-specific helpers:
SUMIFS, CORREL, PERCENTILE, SUMPRODUCT, AVERAGEIFS, INDEX/MATCH.
"""

from openpyxl.utils import get_column_letter


class FormulaBuilder:
    """Utility class for building Excel formula strings."""

    # ── Core Cell References (from IB Toolkit) ─────────────────

    @staticmethod
    def col_letter(col: int) -> str:
        """Convert 1-based column number to letter. col_letter(2) -> 'B'"""
        return get_column_letter(col)

    @staticmethod
    def ref(row: int, col: int, abs_row: bool = False, abs_col: bool = False) -> str:
        """Cell reference: ref(5, 2) -> 'B5', ref(5, 2, True, True) -> '$B$5'"""
        c = get_column_letter(col)
        c_str = f"${c}" if abs_col else c
        r_str = f"${row}" if abs_row else str(row)
        return f"{c_str}{r_str}"

    @staticmethod
    def sheet_ref(sheet: str, row: int, col: int,
                  abs_row: bool = False, abs_col: bool = False) -> str:
        """Cross-sheet reference: sheet_ref('Market_Data', 10, 3) -> \"'Market_Data'!C10\""""
        c = get_column_letter(col)
        c_str = f"${c}" if abs_col else c
        r_str = f"${row}" if abs_row else str(row)
        return f"'{sheet}'!{c_str}{r_str}"

    @staticmethod
    def range_ref(start_row: int, start_col: int, end_row: int, end_col: int,
                  abs_row: bool = False, abs_col: bool = False) -> str:
        """Range reference: range_ref(2, 2, 100, 2) -> 'B2:B100'"""
        sc = get_column_letter(start_col)
        ec = get_column_letter(end_col)
        if abs_col:
            sc, ec = f"${sc}", f"${ec}"
        sr = f"${start_row}" if abs_row else str(start_row)
        er = f"${end_row}" if abs_row else str(end_row)
        return f"{sc}{sr}:{ec}{er}"

    @staticmethod
    def sheet_range(sheet: str, start_row: int, start_col: int,
                    end_row: int, end_col: int,
                    abs_row: bool = False, abs_col: bool = False) -> str:
        """Cross-sheet range: sheet_range('Market_Data', 2, 3, 253, 3) -> \"'Market_Data'!C2:C253\""""
        sc = get_column_letter(start_col)
        ec = get_column_letter(end_col)
        if abs_col:
            sc, ec = f"${sc}", f"${ec}"
        sr = f"${start_row}" if abs_row else str(start_row)
        er = f"${end_row}" if abs_row else str(end_row)
        return f"'{sheet}'!{sc}{sr}:{ec}{er}"

    # ── Basic Formulas (from IB Toolkit) ───────────────────────

    @staticmethod
    def sum_range(start_row: int, end_row: int, col: int) -> str:
        """SUM formula: sum_range(4, 6, 2) -> '=SUM(B4:B6)'"""
        c = get_column_letter(col)
        return f"=SUM({c}{start_row}:{c}{end_row})"

    @staticmethod
    def sum_row_range(row: int, start_col: int, end_col: int) -> str:
        """SUM across columns: sum_row_range(18, 2, 6) -> '=SUM(B18:F18)'"""
        sc = get_column_letter(start_col)
        ec = get_column_letter(end_col)
        return f"=SUM({sc}{row}:{ec}{row})"

    @staticmethod
    def pct(numerator_row: int, denominator_row: int, col: int) -> str:
        """Division: pct(13, 10, 3) -> '=C13/C10'"""
        c = get_column_letter(col)
        return f"={c}{numerator_row}/{c}{denominator_row}"

    @staticmethod
    def iferror(formula: str, fallback: str = "0") -> str:
        """Wrap in IFERROR. Strips leading '=' from inner formula."""
        inner = formula.lstrip("=")
        return f"=IFERROR({inner},{fallback})"

    @staticmethod
    def max_zero(formula: str) -> str:
        """Wrap in MAX(0,...). Strips leading '='."""
        inner = formula.lstrip("=")
        return f"=MAX(0,{inner})"

    # ── Portfolio-Specific Formulas ────────────────────────────

    @staticmethod
    def sumifs(sum_range: str, criteria_range: str, criteria: str) -> str:
        """SUMIFS formula for sector/geographic aggregation.
        sumifs('D2:D50', 'E2:E50', '\"Technology\"') -> '=SUMIFS(D2:D50,E2:E50,\"Technology\")'
        """
        return f"=SUMIFS({sum_range},{criteria_range},{criteria})"

    @staticmethod
    def correl(range1: str, range2: str) -> str:
        """CORREL formula for correlation matrix.
        correl('C2:C253', 'D2:D253') -> '=CORREL(C2:C253,D2:D253)'
        """
        return f"=CORREL({range1},{range2})"

    @staticmethod
    def percentile(data_range: str, k: float) -> str:
        """PERCENTILE.INC formula for VaR.
        percentile('F2:F253', 0.05) -> '=PERCENTILE.INC(F2:F253,0.05)'
        """
        return f"=PERCENTILE.INC({data_range},{k})"

    @staticmethod
    def sumproduct(range1: str, range2: str) -> str:
        """SUMPRODUCT formula for portfolio value calculation.
        sumproduct('B2:Z2', 'B3:Z3') -> '=SUMPRODUCT(B2:Z2,B3:Z3)'
        """
        return f"=SUMPRODUCT({range1},{range2})"

    @staticmethod
    def averageifs(avg_range: str, criteria_range: str, criteria: str) -> str:
        """AVERAGEIFS formula for CVaR.
        averageifs('F2:F253', 'F2:F253', '\"<=\"&G2') -> '=AVERAGEIFS(F2:F253,F2:F253,\"<=\"&G2)'
        """
        return f"=AVERAGEIFS({avg_range},{criteria_range},{criteria})"

    @staticmethod
    def index_match(return_range: str, lookup_range: str, lookup_value: str) -> str:
        """INDEX/MATCH formula for cross-sheet lookups.
        index_match('B2:B50', 'A2:A50', '\"AAPL\"') -> '=INDEX(B2:B50,MATCH(\"AAPL\",A2:A50,0))'
        """
        return f"=INDEX({return_range},MATCH({lookup_value},{lookup_range},0))"

    @staticmethod
    def stdev(data_range: str) -> str:
        """STDEV.S formula: stdev('F2:F253') -> '=STDEV.S(F2:F253)'"""
        return f"=STDEV.S({data_range})"

    @staticmethod
    def count(data_range: str) -> str:
        """COUNT formula: count('F2:F253') -> '=COUNT(F2:F253)'"""
        return f"=COUNT({data_range})"

    @staticmethod
    def average(data_range: str) -> str:
        """AVERAGE formula: average('F2:F253') -> '=AVERAGE(F2:F253)'"""
        return f"=AVERAGE({data_range})"

    @staticmethod
    def max_val(data_range: str) -> str:
        """MAX formula: max_val('F2:F253') -> '=MAX(F2:F253)'"""
        return f"=MAX({data_range})"

    @staticmethod
    def min_val(data_range: str) -> str:
        """MIN formula: min_val('F2:F253') -> '=MIN(F2:F253)'"""
        return f"=MIN({data_range})"

    @staticmethod
    def countifs(criteria_range: str, criteria: str) -> str:
        """COUNTIFS formula: countifs('E2:E50', '\"Technology\"') -> '=COUNTIFS(E2:E50,\"Technology\")'"""
        return f"=COUNTIFS({criteria_range},{criteria})"

    @staticmethod
    def sqrt(formula: str) -> str:
        """SQRT wrapper: sqrt('252') -> '=SQRT(252)'"""
        inner = formula.lstrip("=")
        return f"=SQRT({inner})"
